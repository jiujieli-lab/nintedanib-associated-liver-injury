from pathlib import Path
import json, re, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/workspace/scratch/c5634a2d588f/work/jcm_revision')
OUT=ROOT/'submission/Supplementary_Materials'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.85);s.left_margin=s.right_margin=Inches(.9)
s.header_distance=s.footer_distance=Inches(.35)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman')
 st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman')
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  st._element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
normal=doc.styles['Normal'];normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(10)
for h,size,before,after in [('Heading 1',13,12,7),('Heading 2',11.5,9,5)]:
 st=doc.styles[h];st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
 st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(10.5)
doc.styles['Caption'].paragraph_format.space_after=Pt(7)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(10)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='Additional file 7 Screening methods and results'
doc.core_properties.subject='PDGFRA and FLT4 focused small molecule screening'
doc.core_properties.author='';doc.core_properties.last_modified_by=''

def p(text,style=None):
 z=doc.add_paragraph(text,style);return z
def h(text,level=1):return doc.add_heading(text,level)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,label in enumerate(headers):t.rows[0].cells[i].text=label
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    q=OxmlElement('w:'+edge);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');borders.append(q)
   tcpr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    q=OxmlElement('w:'+edge);q.set(qn('w:w'),'90');q.set(qn('w:type'),'dxa');mar.append(q)
   tcpr.append(mar)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');tcpr.append(shade)
   for para in cell.paragraphs:
    if len(rows)<=4 and ri<len(t.rows)-1:para.paragraph_format.keep_with_next=True
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.04
    for run in para.runs:run.font.size=Pt(10.5);run.font.bold=(ri==0)
 p('')
 return t

p('Additional file 7', 'Title')
p('Screening methods and results', 'Title')
p('Focused comparison of PDGFRA and FLT4 engagement by nintedanib and mechanistic or clinical reference compounds')
p('This supplement describes the chemical identities, experimental pharmacology, computational protocols and validation of a 16-compound comparison following PDGFRA and FLT4 prioritization. The analyses distinguish target engagement from the biological direction of an intervention. They identify testable receptor and ligand hypotheses; they do not establish hepatoprotection, clinical substitution, or reduced drug-induced liver injury (DILI). Detailed source tables are provided in Additional file 8, Tables N20–N27, and the reproducibility data in Additional file 9.')

h('S1 Study design and biological interpretation')
p('The focused library was designed to compare differential engagement of PDGFRA and FLT4, rather than maximize simultaneous inhibition of both receptors. PDGFRA and FLT4 were treated as prioritized mechanistic targets whose causal roles and intervention directions require experimental resolution. A receptor implicated in an injury network may mediate injury, adaptation, repair, or several context-dependent responses. Binding predictions therefore cannot assign a beneficial sign to receptor perturbation.')
p('There is a biological reason to retain competing hypotheses. Hepatic stellate cell-specific Pdgfra deletion reduced fibrosis and supported repair in a carbon tetrachloride injury model, but the effects varied with injury model and time [81]. Conversely, VEGFR3/FLT4 inhibition impaired lymphangiogenesis, lymphatic drainage and reparative macrophage accumulation after hepatic ischemia–reperfusion, while VEGF-D supported repair [82]. These results motivate cell-specific testing of pathological PDGFRA signaling and preservation of FLT4-dependent repair. They do not demonstrate that either receptor mediates nintedanib-associated DILI or that systemic dual inhibition is hepatoprotective.')
p('Nintedanib was the index-drug reference. Imatinib, crenolanib, avapritinib and ripretinib supplied PDGFRA-directed comparator chemistry; axitinib, tivozanib, lenvatinib, sunitinib, pazopanib, sorafenib, regorafenib and ponatinib broadened VEGFR and multikinase engagement patterns. MAZ51 was a nonclinical FLT4 research comparator. Pirfenidone and nerandomilast were clinical nonkinase references, not experimentally established nonbinders to either receptor. Crenolanib remained investigational and MAZ51 a research compound; the other compounds had human regulatory indications at the source review date. Oncology approval was not treated as evidence of suitability for liver disease.')
p('The same 16 parent identities were used for the two receptor screens, the physicochemical calculations and the primary docking campaign. The analysis combined four distinct evidence types: endpoint-separated biochemical records, BoltzMol model outputs, seeded rigid-receptor docking and Inductive Bio physicochemical predictions. These evidence types were reported separately. No composite efficacy or liver-safety score was constructed, and missing experimental or predicted values were not replaced by zeros.')
table(['Analysis','Scope','Principal output'],[
 ['Experimental pharmacology','16 compounds × 2 targets','191 retrieved activity records; 62 eligible exact records, including 56 biochemical records'],
 ['BoltzMol screening','32 submitted compound–target pairs','24 scored structures; 8 filtered dispositions retained'],
 ['Primary docking','16 compounds × 2 receptors × 3 seeds','96 searches and 955 retained poses'],
 ['Separate docking checks','3 apo-PDGFRA searches and 1 preflight redocking','23 sensitivity poses and 10 control poses'],
 ['Physicochemical prediction','16 compounds × 3 properties','48 returned predictions with model bounds and flags']
],[1.55,2.05,3.1])

h('S2 Chemical identities and clinical provenance')
h('S2 1 Parent structures and stereochemistry',2)
p('Compound structures were reconciled against ChEMBL and PubChem identifiers and the relevant regulatory descriptions. ChEMBL parent structures and source activity records were retained with their identifiers; the nintedanib records correspond to [51,68], and the expanded library is individually documented in Table N20. Exact SMILES, stereospecific InChIKeys, parent-versus-salt decisions, comparator roles and provenance are reported in Table N20 and the chemical identity files in Additional file 9. Salt counterions were not represented as separate docking ligands. Canonicalization was used to compare equivalent descriptors; a name match alone was insufficient to establish identity.')
p('Three identity decisions are particularly relevant. Nintedanib used CHEMBL502835 in the Z keto-enamine form, matched to PubChem CID 9809715 and InChIKey XZXHXSATPCNXJR-ZIADKAODSA-N. A name-based PubChem result representing an alternative tautomer was not substituted. MAZ51 used the specified (3Z) parent corresponding to CHEMBL597366 and PubChem CID 9839842; other isomer records returned by a name search were not pooled. Nerandomilast used the (5R) sulfoxide represented by PubChem CID 166177189 and InChIKey UHYCLWAANUGUMN-SSEXGKCCSA-N, consistent with its prescribing information. The ChEMBL record lacking sulfur stereochemistry was not used as the final stereochemical definition.')
p('These inputs define individual parent states, not a pH-dependent tautomer or protonation ensemble. Hydrogens were added during docking preparation, but alternative microspecies were not enumerated. This distinction is material for ionizable kinase ligands, including nintedanib. Model-returned descriptor equivalence was checked independently of coordinate geometry. Matching a SMILES or InChIKey does not by itself verify chirality or double-bond configuration in a predicted three-dimensional structure.')
h('S2 2 Regulatory and hepatic context',2)
p('The regulatory review separated approved indication, molecular mechanism, target-specific experimental evidence, hepatic adverse-reaction warnings and restrictions associated with hepatic impairment. Label wording such as PDGFR or VEGFR inhibition was not converted into an isoform-specific potency value. FLT4 was consistently mapped to VEGFR3 and was not confused with FLT3. Source label URLs and review dates are retained in Table N20 and the label provenance records.')
p('Hepatic liabilities constrain interpretation of the panel. The reviewed US labels for sunitinib, pazopanib, regorafenib and ponatinib contain boxed hepatotoxicity warnings. Nintedanib and pirfenidone also carry DILI warnings. Hepatic impairment can alter nintedanib exposure [14–16]. Other compounds have different warning and exposure profiles, but the absence of a dedicated warning does not establish comparative hepatic safety. Nerandomilast provides a distinct pulmonary treatment mechanism; neither its inclusion nor its predicted properties demonstrate head-to-head liver-safety superiority. All oncology kinase inhibitors were retained as mechanistic comparators, without a therapeutic recommendation.')

h('S3 Experimental target pharmacology')
h('S3 1 Retrieval and eligibility',2)
p('ChEMBL activity data were retrieved for the exact library molecules and parent mappings against the human target records CHEMBL2007 (PDGFRA) and CHEMBL1955 (FLT4), with assay and document metadata inspected on 2 October 2026. The retrieval comprised 191 activity records linked to 78 documents. Table N25 preserves each record, source identifier, endpoint, relation, units, assay context, target-species provenance, variant annotation, duplicate flag and final disposition. Retention of an excluded record does not imply acceptance as quantitative support.')
p('The primary quantitative summaries required a numerical value standardized to nM, an exact equality relation, and an endpoint of Kd, Ki or IC50. The endpoints were analyzed separately. Records describing mutations or fusions were excluded from the nonvariant summaries. Potential duplicates, records with validity flags and clearly identified secondary-source reviews were also excluded. Censored relations such as greater than 10,000 nM, nonstandard endpoints and nonquantitative measurements were retained with their original meaning but not entered as exact values. A missing eligible record was interpreted as a gap in evidence, not inactivity.')
p('Human target mapping and experimentally documented wild-type provenance were distinguished. The absence of a variant description does not prove that an assay used an explicitly verified wild-type human construct. Among the 62 eligible exact records, six explicitly described wild type and 56 did not describe a variant. Fifteen specified human origin in the assay description, 23 relied on ChEMBL human-target mapping, and 24 had unknown origin in the description. These provenance fields remain visible so that stricter subsets can be examined without changing the underlying records.')
p('No assay-confidence score threshold was applied because that field was not retrieved in this dataset. The 62 retained records span 11 compounds and 33 documents. Eighteen records had manually verified primary-source status, 43 had not undergone complete manual source review, and one was patent-derived. Eligibility therefore denotes compliance with the stated metadata rules, not independent verification of every underlying study. These source-status limits remain part of the quantitative interpretation.')
p('Assay context was reviewed from the experimental description rather than assigned solely from a database ontology label. Five PDGFRA IC50 records described extracted recombinant protein produced in Sf9 cells with an HTRF readout, despite a cellular BioAssay Ontology classification. They were interpreted as biochemical recombinant-protein assays, with the conflicting raw annotation preserved. Of the 62 eligible exact records, 56 were biochemical, two cellular and four had uncertain assay context. Table N26 separates these categories and preserves all three endpoint types.')
h('S3 2 Quantitative reference measurements',2)
p('The 56 biochemical records comprised 32 for PDGFRA (six Kd, one Ki and 25 IC50 records) and 24 for FLT4 (seven Kd, one Ki and 16 IC50 records). Counts represent retained records under the stated rules, not independent laboratory replications. Assay format, ATP conditions, construct, substrate and source study can affect the numerical endpoint; consequently, neither IC50 and Kd values nor heterogeneous assay records were pooled into a single potency estimate.')
p('A matched-panel Kd comparison was retained separately in Table N27 [73,83]. Nintedanib had PDGFRA and FLT4 Kd values of 16 and 95 nM, respectively, in the Davis panel. Axitinib had corresponding values of 0.51 and 170 nM. Imatinib provided a useful differentiated engagement reference: PDGFRA Kd 31 nM and FLT4 Kd greater than 10,000 nM in the earlier Karaman panel. The imatinib values reproduced in Davis carry potential-duplicate flags and were retained for matched-panel context, not counted as independent confirmatory evidence. The censored FLT4 result was not changed to an exact 10,000 nM measurement.')
table(['Compound','PDGFRA Kd nM','FLT4 Kd nM','Interpretation'],[
 ['Nintedanib','16','95','Davis matched biochemical reference'],
 ['Axitinib','0.51','170','Davis matched biochemical reference'],
 ['Imatinib','31','>10,000','Karaman result repeated with duplicate flags in Davis']
],[1.3,1.15,1.2,3.05])
p('These reference measurements establish biochemical engagement under their assay conditions; they do not specify receptor occupancy in liver, clinical dosing, or DILI effects. Additional matched-panel entries for sunitinib, pazopanib and sorafenib are preserved in Table N27 with duplicate flags. The table is a contextual benchmark and is not an expansion of the independently eligible exact-value dataset.')
p('A source-level numerical inconsistency was retained for crenolanib. The primary 2012 study assigns PDGFRA Kd 3.2 nM in the Results text but 2.1 nM in Table 1, with the PDGFRB values interchanged (https://doi.org/10.1158/1078-0432.CCR-12-0625). The evidence supports PDGFRA inhibition, but a single definitive Kd was not assigned for calibration. Its directly retrieved parent-molecule ChEMBL record derived from a later review and was excluded from the exact-value summary. MAZ51 was likewise not assigned a numerical IC50 from a single-concentration selectivity observation (https://doi.org/10.1046/j.1432-1033.2001.02476.x).')

h('S4 BoltzMol screening and structure assessment')
h('S4 1 Inputs and model outputs',2)
p('The two focused screens used the Boltz small-molecule library interface, boltz-api version 0.43.0, with the provider-reported boltzmol engine and pipeline version 1.0 [84]. Each job contained all 16 exact parent structures and one protein chain A. The PDGFRA input was the canonical P16234 kinase-domain sequence spanning residues 593–954 (362 residues); FLT4 used P35916 residues 845–1173 (329 residues). Nintedanib was supplied as an input reference ligand in both jobs. Exact payloads, sequence strings, response metadata and structure hashes are archived in Additional file 9.')
p('Provider-default filters were retained. The predefined ranking variable was binding_confidence within each target; optimization_score and global structure-confidence quantities were preserved separately. No threshold was interpreted as a calibrated probability of biochemical binding or a clinical response rate. The optimization metric was not relabeled as Kd, IC50, a binding free energy or an experimental pKd. Model scores from the two targets were displayed separately, without claiming calibrated cross-target selectivity. Provider physicochemical annotations were also retained as separate outputs rather than merged with Inductive Bio predictions.')
h('S4 2 Screening results and filtering',2)
p('Each target job returned scores and coordinates for 12 compounds. Nintedanib, sunitinib, ponatinib and MAZ51 were filtered by the default provider structural-alert filter in both jobs; the individual rule was not returned. Thus, 24 of the 32 submitted compound–target pairs were scored and eight retained a filtered disposition (Table N21). No score or structure was imputed for a filtered compound. Structural-alert filtering cannot be interpreted as experimental nonbinding, and the filtered nintedanib reference cannot serve as a numerical baseline for these model outputs.')
p('For PDGFRA, the five highest binding_confidence values were crenolanib 0.9209, ripretinib 0.8489, axitinib 0.8453, tivozanib 0.8215 and imatinib 0.8045. For FLT4, the leading values were tivozanib 0.7887, lenvatinib 0.7541, pazopanib 0.7465, axitinib 0.7429 and ripretinib 0.7369. Imatinib had a lower FLT4 value of 0.2446. This pattern supplies model-level engagement hypotheses and should be considered alongside the endpoint-separated biochemical data.')
p('Pirfenidone and nerandomilast had lower binding_confidence values in both screens: 0.2323 and 0.0275 for PDGFRA and 0.1394 and 0.0353 for FLT4, respectively. These values do not establish receptor inactivity. The two compounds remain clinical nonkinase references, and no experimentally validated negative-control threshold was inferred from their model scores. Every returned and filtered entry is visible in Table N21.')
h('S4 3 Coordinate and chemical quality assessment',2)
p('All 24 returned CIF files were inspected using their actual result-index paths. Declared chain A sequences matched the input constructs, and all observed residue identities agreed with the declared sequences after mapping CIF residue numbers to UniProt (PDGFRA offset +592; FLT4 offset +844). However, coordinate coverage was incomplete. Every PDGFRA export contained 199 of 362 residues, with native residues 695–712, 728–779, 793–799, 855–878 and 893–954 absent. Every FLT4 export contained 195 of 329 residues, lacking 951–973, 985–995 and 1074–1173. The same omission pattern occurred across the 12 ligands for each target.')
p('Peptide C–N distances within retained consecutive segments were consistent with a continuous local backbone, but large separations occurred across the missing internal regions. A complete sequence declaration therefore did not establish coordinate completeness. All atom B fields equaled 100.000, and no local quality-assessment records were present. These values were not used as local pLDDT or for confidence coloring. Global model-confidence quantities were retained separately; local confidence in the exported pocket and missing regions could not be established from those fields.')
p('Ligand element counts, atom names, connectivity endpoints and input descriptors agreed across the 24 structures. No protein–ligand heavy-atom separation below 2 Å or ligand bond above 2 Å was found. This is a limited severe-overlap and bond-length screen, not complete steric, energetic or stereochemical validation. Coordinate stereochemistry was not independently verified. Descriptor matches cannot establish the correctness of three-dimensional chirality, protonation or tautomer choice.')
p('The representative PDGFRA–imatinib pocket contained the expected 37 ligand heavy atoms and 41 bonds. Its closest protein–ligand pair was Thr674 OG1 to ligand N59 at 2.807 Å; Cys677 backbone N to ligand N34 was 2.930 Å. The FLT4–tivozanib pocket contained 32 ligand heavy atoms and 35 bonds, with Glu896 OE2 to ligand N47 at 2.823 Å and Cys930 backbone N to ligand N23 at 2.975 Å. These are heavy-atom proximities, not assigned hydrogen bonds. Structural depictions use explicit ligand bond types, neutral protein coloring and trace breaks at missing residues. The contact archive contains 449 native-numbered residue contacts across all 24 complexes.')

h('S5 Docking methods and validation')
h('S5 1 Receptor structures and search boxes',2)
p('The primary PDGFRA receptor was chain A of the human wild-type imatinib complex 6JOL at 1.90 Å resolution [89]. All 279 coordinate-bearing residues matched canonical P16234, including wild-type Thr674 and Asp842. Water and ligand were removed while deposited protein heavy atoms were retained. Alternative conformer A was preferred. Meeko 0.8.0 templates assigned receptor chemistry and polar hydrogen positions, retaining all deposited protein residues [87]. The receptor contains the crystallographic kinase-insert deletion and unresolved loops. Its imatinib-bound inactive conformation was fixed throughout the primary library screen and may favor compatible ligand conformations.')
p('No experimentally determined FLT4 kinase-domain coordinates were identified in the RCSB search for P35916. The primary FLT4 receptor was AlphaFold DB AF-P35916-F1 version 6, cropped to residues 845–1173 [90–92]. This is a predicted rigid receptor. The experimental VEGFR2/KDR–nintedanib complex 3C7Q defined the ATP-site search region by homology: BLOSUM62 sequence alignment followed by iterative Cα core superposition retained 222 core pairs with 1.074 Å RMSD from 269 initially matched pairs. Transferred nintedanib coordinates defined only the search box, not a predicted FLT4 pose. Modified or phosphorylated protein residues from the template were not transferred.')
p('Each box used the reference-ligand heavy-atom bounds expanded by 6 Å on every face, with a minimum dimension of 22 Å. Box centers and dimensions are listed below. The primary FLT4 pocket had mean pLDDT 85.88 among residues within 6 Å of the transferred reference, with several glycine-loop residues below 70. This confidence describes the predicted scaffold, not experimental validation of the receptor or its ligand pose. A target-matched FLT4 self-redocking control was unavailable.')
table(['Receptor','Box center x y z in Å','Box dimensions x y z in Å'],[
 ['PDGFRA 6JOL','−38.4125, 157.0490, 0.7940','22.0000, 28.7220, 25.8400'],
 ['FLT4 AlphaFold','−11.70297, 3.17197, −11.88550','22.00000, 22.75049, 29.37968'],
 ['PDGFRA 8PQJ sensitivity','26.94732, 4.68286, 21.49816','25.22713, 27.43338, 22.00000']
],[1.4,2.65,2.65])
h('S5 2 Ligand preparation and stochastic searches',2)
p('RDKit 2026.03.6 generated each starting conformer independently using ETKDGv3, followed by MMFF minimization for up to 2,000 iterations [88]. All minimizations converged. The supplied parent tautomer, protonation state and stereochemical descriptor were preserved apart from canonicalization and hydrogen addition. Meeko 0.8.0 produced ligand and receptor PDBQT files [87]. This preparation did not enumerate physiologic microspecies, so protonation and tautomer dependence remain important limitations.')
p('AutoDock Vina 1.2.7 used the Vina scoring function, exhaustiveness 32, at most 10 retained poses, minimum pose separation 1 Å, energy window 5 kcal/mol and default grid spacing 0.375 Å [86]. Seeds 104729, 130363 and 155921 were fixed before the primary campaign. Each search used two CPU threads; the receptor and box for a target were unchanged across all 16 compounds. The resulting design comprised 16 compounds × two receptors × three searches, giving 96 primary searches.')
p('For each compound–target pair, the three top-ranked scores were summarized by their median and observed range. Top-pose variability was quantified with symmetry-aware heavy-atom RMSD in the unchanged receptor coordinate frame, without ligand superposition. This metric reflects pose displacement and orientation as well as internal geometry. It was not obtained from Vina’s reported upper-bound RMSD field. Repeated seeds assess stochastic search consistency; they are not biological replicates or independent dynamical trajectories. No inferential test used the seeds as experimental replicates.')
h('S5 3 Imatinib redocking control',2)
p('Imatinib was removed from 6JOL and independently embedded from the authoritative CCD STI descriptor before redocking. The separate preflight used the stated box and exhaustiveness and yielded a top-pose symmetry-aware heavy-atom RMSD of 0.574 Å against the deposited ligand, with a Vina score of −12.407 kcal/mol. RMSD was calculated in the unchanged receptor frame without a ligand alignment. The preflight used the CCD atom ordering and four CPU threads.')
p('The three imatinib searches in the actual 16-compound campaign provided directly matched controls using the library atom ordering and the same two-thread settings as the other compounds. Their crystallographic RMSDs were 0.556, 0.776 and 0.712 Å. These results establish successful redocking for imatinib in the selected PDGFRA conformation. They do not validate every compound, the FLT4 predicted scaffold, cross-target selectivity or a liver-protective mechanism. Validation was therefore reported with its target and ligand scope intact.')

h('S6 Docking results and receptor sensitivity')
h('S6 1 Primary library results',2)
p('All 96 primary searches yielded usable outputs, and all 32 compound–target groups contained three searches (Tables N22 and N23; Figure S13). The campaign retained 955 poses. All exported SDF poses parsed, had finite scores and agreed with the corresponding raw Vina result arrays. Run-level inputs, prepared conformers, configurations, receptor definitions and pose files are preserved in Additional file 9.')
p('The lowest median PDGFRA scores were obtained for ponatinib (−12.823 kcal/mol), imatinib (−12.145), regorafenib (−11.363), sorafenib (−11.073) and pazopanib (−11.021). For FLT4, the corresponding five were crenolanib (−10.430), ponatinib (−10.224), nintedanib (−10.081), imatinib (−10.025) and avapritinib (−9.980). These are separate receptor-specific score lists. Their order differs from the BoltzMol ranking and cannot be interpreted as a common experimental affinity scale.')
p('Nintedanib had a median PDGFRA score of −8.520 kcal/mol (range −8.613 to −8.502), placing it 14th of 16 within that receptor. Its FLT4 median was −10.081 kcal/mol (range −10.117 to −9.781), placing it third. The maximum pairwise top-pose RMSD was 9.351 Å for PDGFRA and 1.788 Å for FLT4. The narrow PDGFRA score range therefore coexisted with substantial pose heterogeneity; it did not identify a unique, reproducible binding mode.')
p('For seed 104729, the top PDGFRA nintedanib pose occupied the back-pocket region and remained 12.290 Å from Cys677 by minimum heavy-atom distance. Rank 3 was only 0.161 kcal/mol higher in score and approached gatekeeper Thr674 to 3.380 Å, but remained 6.629 Å from Cys677. Gatekeeper proximity must not be described as a canonical hinge contact. The minimum distances to hinge residues 675–677 were 10.741 and 5.266 Å for ranks 1 and 3, respectively. All retained primary nintedanib poses were assessed, rather than selecting only a favorable geometry.')
h('S6 2 Separately specified apo receptor sensitivity',2)
p('A nintedanib-only PDGFRA sensitivity analysis used the wild-type apo structure 8PQJ, chain A, at 1.82 Å resolution [93]. Although the associated study concerns avapritinib-related chemistry, 8PQJ itself is an apo entry. All 329 coordinate-bearing residues matched canonical P16234, including Thr674 and Asp842. Core alignment to 6JOL retained 271 Cα pairs at 0.785 Å RMSD from 279 initially matched pairs. This single alternate structure is a receptor sensitivity check, not an ensemble sampling the full range of kinase conformations.')
p('The sensitivity protocol was specified after observing the primary nintedanib binding-mode ambiguity and before generating the sensitivity scores. It was therefore a follow-up analysis, not a second independently prespecified primary receptor. The 6JOL imatinib coordinates were transferred only to define the 8PQJ search box using the same expansion rule. No score was used to choose the box. Nintedanib was searched with the same three seeds, ligand preparation, exhaustiveness and two-thread setting as the primary campaign. The primary library rankings were not replaced or recalculated from 8PQJ.')
p('Eleven deposited residues in 8PQJ lacked side-chain atoms. PDBFixer 1.12.0 and OpenMM 8.4.0 added 32 side-chain heavy atoms and one terminal OXT while missing-residue reconstruction was disabled. Every deposited atom retained its original coordinates, with maximum displacement 0 Å. Newly added atoms of Asn780 and Asp785 were then rotated about the CA–CB axes in 15° increments to resolve local clashes; deposited atoms were not moved. These two sites were more than 17 Å from the transferred reference ligand. No missing loop was reconstructed and no receptor MD was part of this preparation.')
p('All three 8PQJ searches completed and retained 23 poses. The median score was −6.495 kcal/mol (range −6.884 to −6.481), with a maximum pairwise top-pose RMSD of 5.094 Å. Top-pose minimum Cys677 distances were 4.729, 3.578 and 4.311 Å, while Thr674 remained 8.238–9.320 Å away. The change demonstrates receptor-dependent proximity and scoring, but does not establish a hydrogen bond or a unique canonical binding mode. All six top poses and all 53 retained PDGFRA nintedanib poses from both receptors are included in the sensitivity data (Figure S14).')
table(['Nintedanib analysis','Median Vina score kcal/mol','Observed score range','Maximum top pose RMSD Å'],[
 ['Primary PDGFRA 6JOL','−8.520','−8.613 to −8.502','9.351'],
 ['Primary FLT4 predicted','−10.081','−10.117 to −9.781','1.788'],
 ['Separate PDGFRA 8PQJ','−6.495','−6.884 to −6.481','5.094']
],[2.0,1.6,1.8,1.3])
p('The total executed docking inventory comprised 100 searches and 988 retained poses: 96 primary searches with 955 poses, three apo sensitivity searches with 23 poses, and one separate preflight control with 10 poses. Table N22 contains only the 96 primary searches. The additional apo searches and complete sensitivity pose inventory are separate CSV files in Additional file 9.')

h('S7 Physicochemical predictions')
p('Inductive Bio predictions were generated for the exact 16 parent structures using mcp_public_logd version 1.7.0 and mcp_public_apka and mcp_public_bpka version 1.4.0 [85]. These models estimate logD at pH 7.4 and the most acidic and basic pKa, respectively. Model identifiers, configuration hashes, returned values, categories, model bounds and applicability flags are reported in Table N24. For nintedanib, predictions obtained with identical model and configuration versions were included; each compound–property combination contributed one value.')
p('All 48 requested compound–property values were returned successfully. Stereospecific InChIKey comparison verified all 16 input identities against the returned descriptors. Nerandomilast was canonicalized by the provider, but its stereospecific identity matched the selected R-sulfoxide. Eight predictions were flagged outside the model domain: acidic and basic pKa for crenolanib, avapritinib, ripretinib and nerandomilast. Three acidic pKa predictions were flagged low confidence (ripretinib, sunitinib and pirfenidone), with overlap between the two flag categories for ripretinib.')
p('Nintedanib had predicted logD at pH 7.4 of 2.523, acidic pKa 10.945 and basic pKa 7.712. The corresponding returned bounds were 1.579–3.447, 7.998–11.196 and 7.409–8.766. The provider did not specify statistical coverage for these bounds, so they were not interpreted as confidence intervals. Pirfenidone and nerandomilast had predicted logD values of 1.376 and 2.680, respectively. Values at the returned pKa limits, including 2 or 12, were preserved without treating them as experimentally measured dissociation constants.')
p('The nintedanib basic pKa estimate identifies protonation as a relevant uncertainty in the single-parent-state docking protocol. These calculations did not assign the site-specific microspecies populations used in a binding experiment, measure free hepatic exposure, or predict DILI. Differences from nintedanib in Table N24 are descriptive differences within the same property model. They are not safety margins or evidence that a candidate would maintain pulmonary efficacy. Figure S14 displays all 48 predictions and their flags, without removal of less reliable estimates.')

h('S8 Integrated interpretation and validation priorities')
p('The three computational views were informative but not interchangeable. BoltzMol favored crenolanib for PDGFRA and tivozanib for FLT4 among the 12 scored compounds per target. Vina produced different receptor-specific orders and supplied scores for all 16 compounds, including the four filtered by BoltzMol. The experimental matched Kd panel independently demonstrated differentiated engagement for imatinib and axitinib. Differences between these outputs reflect their distinct models, receptor representations, chemical coverage and endpoints; no method was selected retrospectively because it gave a preferred compound order.')
p('The results support mechanistic follow-up rather than nomination of a hepatoprotective treatment. A PDGFRA-directed comparator with limited FLT4 engagement may be useful for testing whether profibrotic signaling can be separated from lymphatic repair, but that direction remains a hypothesis. Stronger FLT4 inhibition could oppose repair in some liver-injury settings [82]. Conversely, a weak binding prediction is not proof that a receptor is spared at clinically relevant unbound exposure. Labelled hepatic liabilities and pulmonary efficacy must be evaluated independently of computational engagement.')
p('A discriminating validation strategy would measure target phosphorylation and exposure-matched phenotypes in separately interpretable hepatocyte, stellate-cell, endothelial, lymphatic-endothelial and macrophage systems. The broad endothelial identity in a network does not establish a lymphatic-endothelial mechanism. Single and combined genetic perturbation, pharmacological inhibition, receptor re-expression or a validated resistant construct, and pulmonary antifibrotic readouts could distinguish target dependence from off-target toxicity. Ligand-based rescue would require confirmed receptor signaling because continued kinase inhibition may prevent rescue despite a relevant pathway. These are proposed experiments, not completed observations.')
p('No target–ligand production molecular dynamics simulation or trajectory-based binding-stability result is reported. The docking receptors contain missing or uncertain regions, and all returned BoltzMol models are coordinate crops. Simulation would require a justified receptor construct, reconstruction or explicit treatment of missing regions, reviewed protonation and ligand parameters, and independently initialized trajectories with sampling assessment. Energy minimization, a predicted pose, or a short relaxation cannot establish stable binding. Even a persistent pose would not by itself establish receptor inhibition, activation, affinity, hepatoprotection or clinical benefit.')
p('The principal limitations are a focused rather than exhaustive library, incomplete and heterogeneous experimental coverage, provider filtering, predicted or conformation-specific receptors, single parent chemical states, rigid docking and absence of direct biological intervention experiments. In particular, scores from different receptor structures are not calibrated selectivity measurements, and computational repetitions are not biological sample sizes. The complete dispositions and pose alternatives are retained to make these limitations assessable.')

h('S9 Guide to supplementary tables and reproducibility data')
p('Additional file 8 contains the screening workbook. The table numbers below are stable identifiers for the evidence layers; the workbook is the authoritative record of full rows and numerical precision.')
table(['Table','Content','Reading guidance'],[
 ['N20','Chemical library and clinical provenance','Exact parents, identifiers, roles, regulatory context and hepatic liabilities'],
 ['N21','BoltzMol dispositions and model outputs','32 submitted pairs; 24 scored and 8 filtered; separate confidence and optimization fields'],
 ['N22','Primary docking searches','96 searches only; one row per compound, receptor and seed'],
 ['N23','Primary docking summaries','32 groups; three searches each; score medians, ranges and pose variability'],
 ['N24','Inductive Bio property predictions','48 model outputs; versions, bounds, domain/confidence flags and identity checks'],
 ['N25','Complete experimental assay disposition','191 records including excluded, censored, variant and duplicate entries'],
 ['N26','Endpoint separated pharmacology','96 compound–target–endpoint rows; Kd, Ki and IC50 retained separately'],
 ['N27','Matched Davis panel Kd benchmark','12 contextual records for six compounds; censoring and duplicate flags retained']
],[.65,2.1,3.95])
p('Additional file 9 contains the machine-readable source tables, exact input structures and sequences, provider outputs, docking configurations and poses, receptor-preparation audits, contact tables, and scripts. The nintedanib apo sensitivity is recorded separately in PDGFRA_Nintedanib_Receptor_Sensitivity.csv, PDGFRA_Nintedanib_Sensitivity_All_Poses.csv and PDGFRA_8PQJ_Sensitivity_Summary.json. The run inventory separates primary, sensitivity and preflight calculations. Structure hashes and exact software versions support traceability.')
p('For each docking group, the reproducibility data permit recovery of the starting conformer, seed, fixed receptor and box, top-pose score and all retained pose coordinates. For each predicted property, the exact molecule, model and configuration can be recovered. For each experimental value, the activity identifier, assay description, publication and inclusion or exclusion reason remain linked. These records allow model outputs and experimental evidence to be evaluated without changing the selected chemical identities or discarding inconsistent observations.')

h('Supplementary figure legends')
for number in [13,14]:
 text=(ROOT/f'figures/screening/Figure_S{number}_Legend.txt').read_text().strip()
 first,body=text.split('\n',1)
 p(first,'Heading 2');p(body)

h('References').paragraph_format.page_break_before=True
p('Reference numbers correspond to the main manuscript. Database identifiers and source URLs specific to the expanded library are also retained in Tables N20 and N25 and Additional file 9.')
refs=json.loads((ROOT/'manuscript/Final_References.json').read_text())
used=[14,15,16,51,68,73,81,82,83,84,85,86,87,88,89,90,91,92,93]
for ref in refs:
 if ref['number'] in used:
  z=p(f"[{ref['number']}] {ref['text']}")
  z.paragraph_format.space_after=Pt(7);z.paragraph_format.line_spacing=1.04
  for run in z.runs:run.font.size=Pt(10.5)
p('Additional database records: PubChem nintedanib CID 9809715 (https://pubchem.ncbi.nlm.nih.gov/compound/9809715); MAZ51 CID 9839842 (https://pubchem.ncbi.nlm.nih.gov/compound/9839842); nerandomilast CID 166177189 (https://pubchem.ncbi.nlm.nih.gov/compound/166177189). ChEMBL target records: CHEMBL2007 for PDGFRA and CHEMBL1955 for FLT4 (https://www.ebi.ac.uk/chembl/). RCSB coordinate records: https://www.rcsb.org/structure/3C7Q and https://www.rcsb.org/structure/8PQJ. Database and label provenance for the expanded screen was reviewed on 2 October 2026.')

p('Regulatory source URLs for the hepatic-warning and stereochemical statements are listed below; label sections and review dates are recorded in Table N20.')
labels=json.loads((ROOT/'candidate_library/label_review.json').read_text())
for label in labels['compounds']:
 if label['compound'] in ['sunitinib','pazopanib','regorafenib','ponatinib','pirfenidone','nerandomilast']:
  z=p(label['compound'].capitalize()+': '+label['label_url'])
  for run in z.runs:run.font.size=Pt(10)

out=OUT/'Additional_file_7_Screening_Methods_and_Results.docx'
for style in doc.styles:
 for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for paragraph in doc.paragraphs:
 for border in list(paragraph._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(out)
text='\n'.join(z.text for z in doc.paragraphs)
(ROOT/'manuscript/Additional7_Authored_Text.txt').write_text(text)
(ROOT/'manuscript/Additional7_Content_Audit.json').write_text(json.dumps({'docx':str(out),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'body_word_count':len(text.split()),'main_reference_numbers':used,'source_figure_legends':[13,14],'no_target_ligand_MD_reported':True},indent=2)+'\n')
print(out)
print('Words excluding tables:',len(text.split()))
