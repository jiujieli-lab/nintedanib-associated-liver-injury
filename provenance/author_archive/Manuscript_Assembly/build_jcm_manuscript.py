from pathlib import Path
import json,re,copy,shutil
from docx import Document
from docx.shared import Inches,Pt,Mm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission/Manuscript';OUT.mkdir(parents=True,exist_ok=True)
BASE=ROOT/'journal_audit/content_base'
doc=Document(BASE/'JCM_Content_Base.docx')
new=json.loads((ROOT/'manuscript/new_sections.json').read_text())
refs=json.loads((ROOT/'journal_audit/References_JCM_ACS.json').read_text())+json.loads((ROOT/'journal_audit/New_References_JCM_ACS.json').read_text())
refby={r['number']:r for r in refs}
pat_math=re.compile(r'10([−-]\d+)|(?<![A-Za-z])([KI])([di])\b|\b(IC)(50)\b|\b(log)(2|10)\b|\b(CCl)(4)\b')
def put(p,t,style=None):
    p.clear()
    if style:p.style=style
    last=0
    for m in pat_math.finditer(t):
        p.add_run(t[last:m.start()])
        if m[1]:p.add_run('10');p.add_run(m[1]).font.superscript=True
        else:
            g=m.groups();i=next(i for i in (1,3,5,7) if g[i] is not None)
            p.add_run(g[i]);p.add_run(g[i+1]).font.subscript=True
        last=m.end()
    p.add_run(t[last:]);return p
def find(t):
    a=[p for p in doc.paragraphs if p.text==t];assert len(a)==1,(t,len(a));return a[0]
def before(anchor,text,style='Normal'):
    return put(anchor.insert_paragraph_before(),text,style)
def remove(p):p._element.getparent().remove(p._element)
def after(anchor,text='',style='Normal'):
    p=doc.add_paragraph();put(p,text,style);anchor.addnext(p._element);return p

# Front matter: preserve all identities and declarations from the source.
for p,t in zip(doc.paragraphs[8:12],new['abstract']):
    if t.startswith('Results:'):
        t=t.replace('Competitive ranking','FBP1 distinguished drug-related from other acute liver injury (Cliff’s delta −0.505; q=0.000469). Competitive ranking')
    put(p,t)
for p in doc.paragraphs:
    t=p.text
    for old,replacement in [
        ('Supplementary Figures S1–S12','Supplementary Figures S1–S14'),
        ('Figure 5A summarizes the clinically weighted network integration.',''),
        ('The independent evaluation design is shown in Fig. 6A.',''),
        ('These endpoint-resolved findings elaborate the original pharmacology source rather than add an independent validation cohort. ',''),
        ('Detailed effect estimates and covariance sensitivities are reported in Supplementary Results.','Detailed effect estimates and covariance sensitivities are reported in Additional file 4.'),
        ('; Supplementary Results)', '; Additional file 4)')]:t=t.replace(old,replacement)
    if t!=p.text:put(p,t)

anchor=find('2.12. Reproducibility and statistical units')
for n,section in enumerate(new['methods'],12):
    before(anchor,f'2.{n}. '+section['title'],'Heading 2')
    for t in section['paragraphs']:
        t=t.replace('the existing nintedanib predictions used identical model and configuration versions and were retained','all compound predictions used matched model and configuration versions')
        before(anchor,t)
    if n==14:
        before(anchor,'A separately specified receptor-conformation sensitivity evaluated nintedanib against the wild-type apo PDGFRA structure 8PQJ at 1.82 Å resolution [110]. This analysis followed identification of ambiguous nintedanib poses in 6JOL and retained the primary receptor results. The apo structure was aligned to 6JOL to transfer the prespecified pocket, missing heavy atoms were repaired without moving deposited atoms, and the same three seeds and search settings were used. It tested an alternative receptor conformation without replacing the 16-compound primary screen.')
put(anchor,'2.15. Reproducibility and statistical units')

anchor=find('4. Discussion')
before(anchor,'3.8. '+new['results_title'],'Heading 2')
for t in new['results']:
    if t.startswith('Crystallographic redocking'):
        t=t.replace('(Fig. 7D).','(Fig. 7D); all three primary-library searches also reproduced the crystal pose (RMSD, 0.556–0.776 Å).')
    before(anchor,t)
    if t.startswith('Crystallographic redocking'):
        before(anchor,'Nintedanib showed median docking scores of −8.520 kcal/mol for PDGFRA 6JOL and −10.081 kcal/mol for predicted FLT4. Its top-pose consistency differed between receptors (maximum pairwise RMSD, 9.351 and 1.788 Å, respectively). The apo PDGFRA sensitivity changed hinge proximity but retained substantial pose variability (median score, −6.495 kcal/mol; maximum RMSD, 5.094 Å). These results define receptor-dependent structural hypotheses without establishing a unique PDGFRA binding mode or calibrated cross-target affinity (Supplementary Figures S13 and S14).')
discuss=find('The next study should implement donor-replicated hepatic perturbation of PDGFRA and FLT4, with KDR and FGFR1 comparators, and measure receptor activity, candidate secreted mediators, intracellular metabolism, and extracellular FBP1, GSTA1, and OTC together. Reciprocal rescue should follow an exposure-controlled injury effect. Parent nintedanib, BIBF 1202, and relevant downstream metabolites should be quantified to distinguish disposition from cellular tolerance [28, 29]. Prospective clinical sampling can then connect serial injury proteins and pharmacokinetics to adjudicated liver-test changes and pulmonary response. Existing circulating-biomarker studies support the feasibility of serial assessment [75, 76]. New spatial transcriptomic and protein atlases can refine localization and guide selection of the definitive cellular system [77, 78].')
for t in new['discussion']:before(discuss,t)
lim=[p for p in doc.paragraphs if p.text.startswith('Several limitations define')][0]
put(lim,lim.text+' The focused drug library was selected for mechanistic and clinical relevance, rather than an unbiased chemical-space benchmark. AI confidence and rigid docking scores were not calibrated to hepatic free-drug exposure; protein conformation, fixed parent microspecies and cropped predicted coordinates limit binding-mode inference. Clinical efficacy and comparative hepatic safety were not evaluated by these calculations.')

# Completed analyses and matching supplementary filenames.
for p in doc.paragraphs:
    if p.text.startswith('Additional file 2.'):
        put(p,'Additional file 2. Supplementary Figures S1–S14. Additional_file_2_Supplementary_Figures.pdf. Source-specific validation, target-pair network interventions, liver-cell localization, numerical calibration, complete docking comparisons and molecular-property predictions.')
    elif p.text.startswith('Additional file 3.'):
        put(p,'Additional file 3. Supplementary tables. Additional_file_3_Supplementary_Tables.xlsx. Complete source-specific quantitative results. Source analysis identifiers are mapped to final figures in the crosswalk on page 93 of Additional file 1.')
    elif p.text.startswith('Additional file 4.'):
        put(p,'Additional file 4. Targeted analyses and extended Results. Additional_file_4_Targeted_Analyses.docx. Complete biochemical-binding and network-intervention results, network and pulmonary sensitivities, and the Figure S11–S12 legends.')
    elif p.text.startswith('The public datasets analysed'):
        put(p,'The public datasets analysed in this study are accessible through the repositories and source publications identified in Tables 1 and 2. Additional files 1–3 provide source-specific extended methods, Supplementary Figures S1–S14 and quantitative tables. Additional files 4–6 provide targeted biochemical and virtual-intervention methods, results, tables and executable source data. Additional files 7–9 provide focused compound-screening methods, all candidate dispositions, molecular-property and docking outputs, exact structures, model versions and reproducible analysis code. The accompanying author reproducibility archive retains the five original source packages and figure-composition records.')
    elif p.text.startswith('OpenAI ChatGPT with Codex tools assisted'):
        put(p,'OpenAI ChatGPT with Codex tools assisted literature retrieval, analytical code development and review, quantitative figure preparation, manuscript organization, language revision and consistency checking during September–October 2026. OpenAI image generation supported non-data biomedical illustrations in the workflow panels. Quantitative panels used numerical source outputs. BoltzMol, Proto and Inductive Bio were used for the computational procedures specified in Methods, with inputs and resolved versions retained. The authors remain responsible for data provenance, analytical choices, interpretation and the final manuscript.')
    elif p.text.startswith('SH contributed to conceptualization'):
        put(p,(BASE/'Author_Contributions_JCM_Format.txt').read_text().strip())
anchor=find('Author Contributions')
for t in [
 'Additional file 7. Focused compound-screening methods and results. Additional_file_7_Screening_Methods_and_Results.docx. Complete screening and docking protocols, receptor and ligand preparation, structure quality checks, and Supplementary Figure S13–S14 legends.',
 'Additional file 8. Screening tables. Additional_file_8_Screening_Tables.xlsx. Supplementary Tables N20–N27 contain the complete 16-compound library, 32 AI-screening dispositions, 96 primary docking searches, 48 property predictions and all experimental activity records.',
 'Additional file 9. Compound-screening source data and code. Additional_file_9_Screening_Code_and_Data.zip. Exact inputs, raw model responses, predicted pocket coordinates, docking structures and logs, three-seed receptor sensitivity, analysis scripts and checksums.'
]:before(anchor,t)

# Collect the existing complete legends before placing figures in the text.
legend_start=find('Figure legends');ps=doc.paragraphs;start=ps.index(legend_start) if legend_start in ps else next(i for i,p in enumerate(ps) if p.text=='Figure legends')
legend={}
for n in range(1,7):
    p=next(p for p in doc.paragraphs if p.text.startswith(f'Figure {n} '))
    idx=next(i for i,q in enumerate(doc.paragraphs) if q._element is p._element)
    legend[n]=(p.text,doc.paragraphs[idx+1].text)
    remove(doc.paragraphs[idx+1]);remove(p)
remove(legend_start)
f7=(ROOT/'figures/screening/Figure_7_Legend.txt').read_text().strip().split('\n\n',1)
legend[7]=(f7[0].replace('Figure 7:','Figure 7.'),f7[1].replace('pipelineboltzmol','pipeline boltzmol').replace('The chemical identities and model-coordinate provenance were verified.','Input chemical identities and coordinate-file provenance were verified.'))
before(find('References'),'','Normal')

# Move editable tables with their complete captions after the first shared citation.
table_objects=list(doc.tables)
table_titles=[]
for n in (1,2):
    p=next(p for p in doc.paragraphs if p.text.startswith(f'Table {n} '));idx=next(i for i,q in enumerate(doc.paragraphs) if q._element is p._element)
    note=doc.paragraphs[idx+1]
    table_titles.append((p,note))
tables_heading=find('Tables');remove(tables_heading)
anchor=next(p for p in doc.paragraphs if 'Tables 1 and 2 identify' in p.text)._element
for (title,note),table in zip(table_titles,table_objects):
    for elm in [title._element,table._element,note._element]:anchor.addnext(elm);anchor=elm
    title.paragraph_format.keep_with_next=True

# Preserve figure pixel data; embedding uses a readable publication preview.
images={1:ROOT/'figures/final/Current_Figure_1.png',2:ROOT/'figures/repaired_quantitative/Figure_2.png',3:ROOT/'figures/repaired_quantitative/Figure_3.png',4:ROOT.parents[1]/'figure4_repair/Figure_4_Repaired_1200dpi.tiff',5:ROOT/'figures/final/Current_Figure_6.png',6:ROOT/'figures/final/Current_Figure_7.png',7:ROOT/'figures/screening/final/Figure_7_Compound_Comparison.png'}
# ROOT is work/jcm_revision; the workspace is two parents above it.
images[4]=ROOT.parents[1]/'figure4_repair/Figure_4_Repaired_1200dpi.tiff'
for n in range(1,8):
    assert images[n].is_file(),images[n]
    pattern=re.compile(r'(?<!Supplementary )\b(?:Fig\.|Figure) '+str(n)+r'(?!\d)')
    candidates=[p for p in doc.paragraphs if pattern.search(p.text) and not p.text.startswith('Figure ')]
    assert candidates,(n,'no first citation')
    anchor=candidates[0]._element
    p=after(anchor);p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True;p.alignment=1
    with Image.open(images[n]) as im:w,h=im.size
    width=min(6.65,8.25*w/h)
    p.add_run().add_picture(str(images[n]),width=Inches(width))
    title=after(p._element,legend[n][0],'Caption');title.paragraph_format.keep_with_next=True
    caption=after(title._element,legend[n][1],'Caption');caption.paragraph_format.keep_with_next=False
    for q in [title,caption]:
        q.paragraph_format.line_spacing=1.0;q.paragraph_format.space_after=Pt(4)
        for r in q.runs:r.font.size=Pt(9)
    for r in title.runs:r.bold=True

# Reference order follows the final body, including relocated tables and captions.
citation=re.compile(r'\[([0-9,;\s–-]+)\]')
def nums(s):
    out=[]
    for x in re.split(r'[,;]\s*',s):
        if not x.strip():continue
        a=re.split(r'[–-]',x.strip())
        out.extend(range(int(a[0]),int(a[1])+1) if len(a)==2 else [int(a[0])])
    return out
def paragraphs_in_order():
    for el in doc._element.body:
        if el.tag==qn('w:p'):
            p=Paragraph(el,doc)
            if p.text=='References':break
            yield p
        elif el.tag==qn('w:tbl'):
            for pel in el.iter(qn('w:p')):yield Paragraph(pel,doc)
order=[]
for p in paragraphs_in_order():
    for m in citation.finditer(p.text):
        for n in nums(m[1]):
            assert n in refby,(n,p.text)
            if n not in order:order.append(n)
assert set(order)==set(refby),(set(refby)-set(order),set(order)-set(refby))
mapping={old:i+1 for i,old in enumerate(order)}
def renumber_para(p):
    text=p.text
    for m in reversed(list(citation.finditer(text))):
        vals=sorted(set(mapping[x] for x in nums(m[1])))
        replacement='['+', '.join(map(str,vals))+']'
        start,end=m.span();runs=p.runs;positions=[];pos=0
        for r in runs:positions.append((pos,pos+len(r.text)));pos+=len(r.text)
        touched=[i for i,(a,b) in enumerate(positions) if a<end and b>start]
        if not touched:continue
        first,last=touched[0],touched[-1]
        prefix=runs[first].text[:start-positions[first][0]];suffix=runs[last].text[end-positions[last][0]:]
        runs[first].text=prefix+replacement+(suffix if first==last else '')
        for i in touched[1:]:runs[i].text=suffix if i==last else ''
for p in paragraphs_in_order():renumber_para(p)
refhead=find('References');seen=False
for el in list(doc._element.body):
    if el is refhead._element:seen=True;continue
    if seen and el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
for old in order:
    r=refby[old];p=doc.add_paragraph();p.add_run(str(mapping[old])+'. ')
    for s in r['segments']:
        rr=p.add_run(s['text']);rr.bold=s.get('bold',False);rr.italic=s.get('italic',False)
    p.paragraph_format.left_indent=Inches(.26);p.paragraph_format.first_line_indent=Inches(-.26)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.0
    for rr in p.runs:rr.font.size=Pt(9.5)

# Neutral submission typography and editable tables.
for sec in doc.sections:
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(20);sec.bottom_margin=Mm(20);sec.left_margin=Mm(20);sec.right_margin=Mm(20)
    sec.header_distance=Mm(8);sec.footer_distance=Mm(8)
    for p in sec.footer.paragraphs:p.clear()
    p=sec.footer.paragraphs[0];p.alignment=1
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
for st in ['Normal','Normal (Web)','Heading 1','Heading 2','Title','Caption']:
    if st in doc.styles:
        s=doc.styles[st];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.line_spacing=1.15
        if st in ['Normal','Normal (Web)']:s.font.size=Pt(11)
        if st=='Caption':s.font.size=Pt(9)
for p in doc.paragraphs:
    if p.style.name!='Caption' and not re.match(r'^\d+\. [A-Z]',p.text or ''):
        p.paragraph_format.line_spacing=1.15
    p.paragraph_format.widow_control=True
    for r in p.runs:r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
for i,t in enumerate(doc.tables):
    t.autofit=False
    widths=[1.45,1.75,3.45] if i==0 else [1.25,1.75,1.35,2.3]
    for j,c in enumerate(t.columns):c.width=Inches(widths[j])
    mar=t._tbl.tblPr.find(qn('w:tblCellMar'))
    if mar is not None:
        for edge in ['top','bottom']:
            el=mar.find(qn('w:'+edge))
            if el is not None:el.set(qn('w:w'),'35')
    for ri,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j])
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_with_next=(ri==0);p.paragraph_format.keep_together=True
                for r in p.runs:r.font.size=Pt(9);r.font.name='Times New Roman'
    pr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');pr.append(rep)
doc.core_properties.title=doc.paragraphs[0].text;doc.core_properties.subject='Original research article for Journal of Clinical Medicine'
out=OUT/'JCM_Manuscript.docx';doc.save(out)
(ROOT/'manuscript/Final_Reference_Number_Map.json').write_text(json.dumps(mapping,indent=2))
(ROOT/'manuscript/Final_References.json').write_text(json.dumps([dict(refby[o],original_number=o,number=mapping[o]) for o in order],indent=2,ensure_ascii=False))
(ROOT/'manuscript/JCM_Manuscript_Text.txt').write_text('\n\n'.join(p.text for p in doc.paragraphs))
report={'references':len(order),'main_figures':7,'tables':len(doc.tables),'abstract_words':sum(len(p.text.split()) for p in doc.paragraphs[8:12]),'original_author_identities_preserved':True,'production_md_claimed':False,'output':str(out)}
(ROOT/'manuscript/Assembly_Report.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
