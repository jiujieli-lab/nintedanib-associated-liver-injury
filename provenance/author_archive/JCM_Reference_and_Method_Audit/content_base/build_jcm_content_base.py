from pathlib import Path
import re,json,copy,hashlib
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
BASE=Path(__file__).resolve().parent
AUDIT=BASE.parent
SOURCE=Path('/workspace/scratch/c5634a2d588f/deliverables/Manuscript/Revised_Manuscript_Clean.docx')
doc=Document(SOURCE)
ps=list(doc.paragraphs)
orig={i:p.text for i,p in enumerate(ps)}
styles={i:p.style.name for i,p in enumerate(ps)}
changes={}
# The original text is retained below for every paragraph moved to supplementary results.
moved=[88,89,90,92,93,94,95,97,100,111]
repl={
8:'Background/Objectives: Liver injury can interrupt nintedanib treatment for fibrosing interstitial lung disease. We connected human injury proteins with drug pharmacology and liver-cell networks to prioritize receptors for mechanistic investigation.',
9:'Methods: We integrated active-comparator pharmacovigilance, 13 human serum proteins, biochemical and cellular drug responses, and a five-donor liver atlas. Donor-balanced networks, matched-background calibration, six aggregation algorithms, donor omission and weight perturbation evaluated candidate selection. Published binding measurements and single/joint network interventions characterized PDGFRA and FLT4. Independent hepatic profiles from 22 mice and pulmonary datasets assessed tissue context.',
10:'Results: PDGFRA and FLT4 ranked first and second under four aggregation rules and remained among the top five in all 5000 weight perturbations. PDGFRA remained first after donor omission; FLT4 ranked second to fifth. Reported dissociation constants were 16 and 95 nM, respectively. The narrow hepatic reporting odds ratio for nintedanib versus pirfenidone was 3.997 (95% confidence interval 3.226–4.951), with temporal heterogeneity. FBP1 distinguished drug-related from other acute liver injury (Cliff’s delta −0.505; q=0.000469). Competitive ranking and network interventions did not establish significant enrichment. Independent liver analysis identified decreased Kdr expression within the candidate family (log2 difference −0.280; q=0.0386; genome-wide q=0.119), without significant ranking–response association. Pulmonary responses varied by model.',
11:'Conclusions: PDGFRA and FLT4 are reproducible priorities for studying nintedanib-associated liver injury. Their biochemical engagement and cellular context support exposure-controlled perturbation and rescue experiments that measure hepatic injury while preserving pulmonary antifibrotic activity. KDR provides a hepatic pharmacodynamic comparator; the protective intervention direction remains to be established.',
13:'Introduction',20:'Materials and Methods',
37:orig[37].replace('within the frozen ChEMBL extraction','within the ChEMBL extraction').replace('The original eligibility rules remained primary.','The specified eligibility rules defined the primary set.').replace('No inferential pooling across assay types or conventional molecular docking was performed.','Assay types were summarized separately without inferential pooling.'),
85:orig[85].replace('The frozen pharmacology set','The primary pharmacology set'),
52:orig[52].replace('Revision-level stochastic calculations','Stochastic calculations'),
55:orig[55].replace('The extension adopted','We applied'),
87:'Clinical-protein-guided integration prioritizes PDGFRA and FLT4',
88:'Liver-cell localization placed FLT4 in the endothelial compartment and PDGFRA predominantly in the sparsely sampled stellate compartment, while the clinical injury proteins were mainly hepatocellular (Supplementary Fig. S12B–G). Numerical calibration excluded 132 of 210 virtual-knockout runs with noise-scale displacement; the remaining drug-target runs showed no enrichment of the clinical protein phenotype (minimum P=0.352; Supplementary Fig. S12H–K). These findings guided integration of clinical-network proximity with biochemical pharmacology and cellular drug responses.',
92:'The donor-balanced analysis evaluated 50 target–compartment combinations across the fixed 17-target universe (Fig. 5B). Primary computations passed numerical checks, and alternative graph, clinical-weight and donor settings supplied complementary tests of stability. Three proximity associations passed multiplicity correction, but no candidate jointly met significance, rank stability and the 10% same-compartment detection threshold (Fig. 5C,M). The macrophage PDGFRA association was dominated by HPD and arose from only two detected cells; seed decomposition, degree-matched calibration and donor omission therefore informed the choice of cellular systems (Fig. 5D,E). Detailed effect estimates and covariance sensitivities are reported in Supplementary Results.',
96:orig[96].replace('The network feature was then combined with pharmacology, HepG2 response, and STRING association context to choose reproducible experiments','Integration of the clinical-network feature with pharmacology, HepG2 response and STRING association context prioritized PDGFRA and FLT4').replace('under revised median','under median'),
97:'Geometric fusion retained the ranking after component omission more consistently than the historical consensus (mean Spearman correlation 0.833 versus 0.793; mean top-five Jaccard overlap 0.607 versus 0.506; Fig. 5H). Agreement with the omitted feature remained weak, and none of the five integration methods reached competitive q<0.05 (Fig. 5I; Supplementary Results). Candidate stability therefore supports the order of experimental testing without establishing predictive accuracy.',
100:'Single and joint PDGFRA/FLT4 network interventions did not exceed matched-control backgrounds (all nine primary q≥0.529), and joint residuals were nonsignificant. These analyses define within-atlas sensitivity rather than a directional or synergistic receptor relationship (Supplementary Fig. S11; Supplementary Results).',
111:'Cellular and protein measurements further delimited pulmonary target confirmation. GSE151374 showed no FDR-controlled effects on the examined macrophage genes, composition or programs at the pooled-library level (Supplementary Fig. S6). In PXD052594, no protein passed q<0.10, and none of the quantified phosphosites in the examined receptor and downstream signaling tracks passed q<0.05 (Supplementary Fig. S7). Pooling, postrandomization losses and the absence of matched total-protein adjustment constrained these comparisons; complete sample counts, nominal effects and phosphoproteome results are retained in Supplementary Results.',
114:orig[114].replace('Numerical auditing sharpened that sequence by identifying noise-scale virtual-knockout outputs that had favored FGFR1. ',''),
118:'Numerical calibration established which virtual-knockout outputs could enter biological interpretation. Rank transformation can order numerical noise; retaining raw distances and applying scale checks prevented those outputs from determining receptor priority. The calibrated network framework then linked candidate selection to the clinical phenotype on an interpretable scale.',
119:orig[119].replace('The revised network analysis','The donor-balanced network analysis'),
120:'Algorithm comparison showed that the leading receptor priorities persisted across specified changes in aggregation, source weighting and donor composition. The weak omitted-feature and independent hepatic rank associations distinguish this reproducibility from predictive validation. Together, these analyses support a focused experimental order whose biological value can be tested directly.',
121:orig[121].replace('The virtual-intervention extension then asks a different question—how changing each receptor node alters the clinically weighted network. Its calibrated results support neither a synergistic interaction nor a directional PDGFRA–FLT4 regulatory relationship.','The network interventions assess sensitivity to receptor-node deletion; their nonsignificant matched-background effects leave the biological direction and interaction between PDGFRA and FLT4 unresolved.'),
124:orig[124].replace('the revised experimental priorities','the integrated experimental priorities'),
127:orig[127].replace('the frozen rankings','the fixed rankings').replace('Revision-developed methods were frozen before','Analysis settings were fixed before'),
130:'Abbreviations',
133:'Institutional Review Board Statement',135:'Informed Consent Statement',
136:'Not applicable to this secondary analysis of deidentified public data. Consent procedures are reported by the source studies. The manuscript contains no identifiable individual participant information or images.',
137:'Data Availability Statement',
138:orig[138].replace('the original complete supplementary tables','the complete supplementary tables').replace('the new analysis scripts','analysis scripts').replace('retains the original five','retains five').replace('This documentation distinguishes archived source results, reconstructed auxiliary files and calibrated analyses. ','').replace('Supplementary Figures S1–S11','Supplementary Figures S1–S12'),
140:'Conflicts of Interest',144:'Author Contributions',146:'Acknowledgments',
149:orig[149].replace('during September–October 2026','during September–October 2026'),
150:'Supplementary Materials',
151:orig[151].replace('References in this supplement are numbered independently.','References use the numbering of the main manuscript.'),
152:'Additional file 2. Supplementary Figures S1–S12. Additional_file_2_Supplementary_Figures.pdf. Supplementary figures covering source-specific validation, target-pair network sensitivity, liver-cell localization and numerical perturbation calibration.',
153:orig[153].replace('preserved without numerical or structural changes','containing the source-specific quantitative results'),
156:orig[156].replace('frozen specifications','analysis specifications'),
165:orig[165].replace('the frozen primary set','the primary set').replace('These are reanalysed experimental records, not newly performed binding assays, docking scores or molecular-dynamics results.','Values are derived from published experimental assays.'),
170:orig[170].replace('the frozen candidate priorities','the fixed candidate priorities'),
}
# Remapping is simultaneous to avoid converting old 6 -> new 5 -> S12.
figpat=re.compile(r'\b(Fig\.|Figure) ([567])(?=[A-Z\b\s.,;–])')
def figmap(t):
 return figpat.sub(lambda m: ('Supplementary Fig. S12' if m[2]=='5' else m[1]+' '+{'6':'5','7':'6'}[m[2]]),t)
def clean(t):
 for a,b in [('Supplementary Figures S1–S11','Supplementary Figures S1–S12'),('frozen ranking','fixed ranking'),('frozen priorities','fixed priorities'),('frozen-algorithm','fixed-algorithm'),('Frozen-algorithm','Fixed-algorithm'),('frozen openFDA','openFDA'),('revised four-feature','four-feature'),('Revised values','Feature values'),('within each revised algorithm','within each integration algorithm'),('other revised rules','other integration rules'),('identical revised inputs','identical four-feature inputs')]:t=t.replace(a,b)
 return t
# New text already using new Figure 5 is excluded from old-number remapping.
new_numbered={92,97}
# Superscript only exact mathematical exponent syntax, and subscript only conventional biochemical labels.
pat=re.compile(r'10([−-]\d+)|(?<![A-Za-z])([KI])([di])\b|\b(IC)(50)\b|\b(log)(2|10)\b|\b(CCl)(4)\b')
def put(p,text,style=None):
 p.clear()
 if style:p.style=style
 pos=0
 for m in pat.finditer(text):
  p.add_run(text[pos:m.start()])
  if m[1]:p.add_run('10');r=p.add_run(m[1]);r.font.superscript=True
  else:
   g=m.groups();a=next(i for i in (1,3,5,7) if g[i] is not None);p.add_run(g[a]);r=p.add_run(g[a+1]);r.font.subscript=True
  pos=m.end()
 p.add_run(text[pos:])
 return p
for i,p in enumerate(ps):
 if i>=178:continue
 t=repl.get(i,orig[i]);t=clean(t)
 if i not in new_numbered:t=figmap(t)
 # Logical zero-power in source already includes unicode superscripts; preserve it.
 if t!=orig[i] or i==145:
  put(p,t,'Normal' if i==145 else None)
  changes[i]={'action':'edited','before':orig[i],'after':t}
# Move whole original paragraphs to supplement, using abbreviated replacements where specified.
remove=[89,90,91,93,94,95,99,132,148,166,167]
for i in remove:ps[i]._element.getparent().remove(ps[i]._element)
# Priority Results begin with integrated ranking, stability, then cellular context and calibration.
anchor=ps[87]._element
for i in [96,98,88,92,97,100]:
 anchor.addnext(ps[i]._element);anchor=ps[i]._element
# Add Methods disclosure consistent with factual acknowledgement, no model-version invention.
ai=ps[69].insert_paragraph_before('OpenAI ChatGPT with Codex tools supported literature retrieval, analytical code development and checking, visualization and manuscript preparation. Source-linked data, numerical consistency checks and author review governed the use of generated outputs. Inductive Bio model versions and inputs are specified above; further details are provided in the Acknowledgments.')
ai.style='Normal'
# Standard JCM backmatter sequence, with exact author/funding/ethics source facts retained.
anchor=ps[129]._element
for i in [150,151,152,153,154,155,156,144,145,142,143,133,134,135,136,137,138,139,146,147,149,140,141,130,131]:
 anchor.addnext(ps[i]._element);anchor=ps[i]._element
# Authors and affiliations retain source spelling and exact order, with actual superscript references.
p=ps[1];p.clear()
for j,(name,aff) in enumerate([('Shaoxin Huang','1,2,3'),('Yong Yang','3'),('Xinyu Zhou','4'),('Xinyi Chen','1'),('Shan Li','1'),('Wenyan Fan','1'),('Jianjun Xiong','2'),('Hui Liu','1,2*')]):
 if j:p.add_run(', ')
 p.add_run(name);p.add_run(aff).font.superscript=True
for i in range(2,6):
 p=ps[i];p.clear();p.add_run(orig[i][0]).font.superscript=True;p.add_run(orig[i][1:])
# Format verified bibliography without changing source reference numbers.
refs=json.loads((AUDIT/'References_JCM_ACS.json').read_text())
for r,p in zip(refs,ps[178:]):
 p.clear();p.add_run(str(r['number'])+'. ')
 for seg in r['segments']:
  run=p.add_run(seg['text']);run.bold=seg.get('bold',False);run.italic=seg.get('italic',False)
# JCM conventional section numbering, while retaining headings of front/back matter.
section=0;sub=0;inmain=False
for p in doc.paragraphs:
 t=p.text
 if p.style.name=='Heading 1' and t in ['Introduction','Materials and Methods','Results','Discussion','Conclusions']:
  section+=1;sub=0;inmain=True;put(p,f'{section}. {t}')
 elif p.style.name=='Heading 1':inmain=False
 elif inmain and p.style.name=='Heading 2':
  sub+=1;put(p,f'{section}.{sub}. {t}')
# Backmatter headings must not inherit numbered Methods/Results styling; preserve declared source facts.
for i in [133,135,137,140,142,144,146]:ps[i].style='Heading 1'
# Consistent black manuscript typography; no numerical character is superscripted merely for being a digit.
for sty in ['Normal','Normal (Web)','Title','Heading 1','Heading 2']:
 if sty not in doc.styles:continue
 s=doc.styles[sty];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
 if sty in ['Normal','Normal (Web)']:s.font.size=Pt(11);s.paragraph_format.space_after=Pt(6)
 if sty=='Title':s.font.size=Pt(16)
 if sty=='Heading 1':s.font.size=Pt(12);s.font.bold=True
 if sty=='Heading 2':s.font.size=Pt(11);s.font.bold=True
for p in doc.paragraphs:
 for r in p.runs:
  r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
# Source tables remain editable and unchanged in sequence; root will place near first citations.
out=BASE/'JCM_Content_Base.docx';doc.save(out)
# Preserve moved material twice: exact source text for disposition and remapped publication text for integration.
suppgroups=[('Liver-cell localization and numerical calibration',[88,89,90]),('Network decomposition and sensitivity analyses',[92,93,94,95,97]),('Single and joint PDGFRA and FLT4 network interventions',[100]),('Pulmonary cellular and protein context',[111])]
parts=[]
for h,ids in suppgroups:
 parts.append(h+'\n')
 for i in ids:parts.append(clean(figmap(orig[i]))+'\n')
parts.append('Supplementary Figure S12. Liver localization and numerical auditing calibrate virtual perturbation evidence\n')
parts.append(orig[167]+'\n')
(BASE/'Supplementary_Results_Integration.txt').write_text('\n'.join(parts))
(BASE/'Moved_Original_Results_Exact.txt').write_text('\n\n'.join(f'Source paragraph {i}\n{orig[i]}' for i in moved)+f'\n\nSource Figure 5 legend\n{orig[166]}\n{orig[167]}\n')
# Paragraph-by-paragraph scientific disposition.
disp=[]
for i in range(275):
 if i not in orig:break
 action='retained'
 if i>=178:action='reference metadata corrected and ACS formatted; source numbering retained'
 elif i in moved:action='complete original in supplementary integration text; abbreviated main text' if i in repl else 'complete original moved to supplementary integration text'
 elif i in remove:action='heading or legend removed from main; Figure 5 legend moved to S12' if i in [166,167] else 'redundant heading removed'
 elif i in changes:action='edited for JCM narrative, terminology, or figure mapping'
 if i in [96,98]:action+='; moved to start of priority Results'
 disp.append({'source_paragraph':i,'source_style':styles[i],'source_text':orig[i],'action':action,'main_text':None if i in remove else ps[i].text})
(BASE/'Paragraph_Disposition.json').write_text(json.dumps(disp,ensure_ascii=False,indent=2))
(BASE/'Paragraph_Disposition.tsv').write_text('source_paragraph\taction\n'+'\n'.join(f"{r['source_paragraph']}\t{r['action']}" for r in disp))
words=lambda s:len(s.split())
report={'source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output':str(out),'abstract_words':sum(words(repl[i]) for i in range(8,12)), 'original_abstract_words':sum(words(orig[i]) for i in range(8,12)), 'original_results_words':sum(words(orig[i]) for i in range(71,113)), 'main_results_words':sum(words(p.text) for p in doc.paragraphs if False), 'source_tables':len(doc.tables), 'reference_count':len(refs), 'figure_mapping':{'1':'1','2':'2','3':'3','4':'4','5':'S12','6':'5','7':'6'},'reserved_main_figure':7,'moved_source_paragraphs':moved,'references_globally_renumbered':False,'figures_embedded':False,'render_status':'Intermediate base; root will render and inspect after final result, figure and reference integration.'}
inside=False;resulttexts=[]
for p in doc.paragraphs:
 if p.text=='3. Results':inside=True;continue
 if p.text=='4. Discussion':inside=False
 if inside:resulttexts.append(p.text)
report['main_results_words']=words(' '.join(resulttexts))
(BASE/'Content_Base_Report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
(BASE/'JCM_Content_Base_Text.txt').write_text('\n\n'.join(p.text for p in doc.paragraphs))
print(json.dumps(report,indent=2))
