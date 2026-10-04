from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import json
B=Path(__file__).resolve().parent;R=B.parent;OUT=R/'submission';(OUT/'Author_Information').mkdir(exist_ok=True)
def newdoc():
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.8)
 for name,sz in [('Normal',11),('Title',16),('Heading 1',12),('Heading 2',11)]:
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6)
  if name=='Normal':st.paragraph_format.line_spacing=1.1
  else:st.font.bold=True;st.paragraph_format.space_before=Pt(8)
 for x in d.styles.element.xpath('.//w:rFonts'):
  for a in list(x.attrib):
   if 'theme' in a.lower():del x.attrib[a]
 for x in d.styles.element.xpath('.//w:pBdr'):x.getparent().remove(x)
 return d
def paragraph(d,t,style=None):
 p=d.add_paragraph(t,style);p.paragraph_format.widow_control=True;return p
def link(d,label,url):
 p=d.add_paragraph();p.add_run(label+': ');h=OxmlElement('w:hyperlink');h.set(qn('r:id'),d.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)
def table(d,rows,widths=(1.65,5.2)):
 t=d.add_table(rows=0,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for j,data in enumerate(rows):
  cells=t.add_row().cells
  for i,txt in enumerate(data):
   cells[i].text=txt;cells[i].width=Inches(widths[i]);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cells[i]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for a in ['top','left','bottom','right']:
    x=OxmlElement('w:'+a);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b);m=OxmlElement('w:tcMar')
   for a,v in [('top','70'),('bottom','70'),('left','100'),('right','100')]:x=OxmlElement('w:'+a);x.set(qn('w:w'),v);x.set(qn('w:type'),'dxa');m.append(x)
   pr.append(m)
   if j==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF4');pr.append(sh)
   for p in cells[i].paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(10);r.bold=j==0
  trpr=t.rows[j]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if j==0:x=OxmlElement('w:tblHeader');trpr.append(x)
 return t
# Clean scientific letter. Required author-controlled declarations remain in the author checklist.
d=newdoc();paragraph(d,'Submission to Journal of Clinical Medicine','Title')
text=(B/'content_base/Cover_Letter.txt').read_text();blocks=text.strip().split('\n\n')
for b in blocks:paragraph(d,b)
letter=OUT/'Cover_Letter.docx';d.save(letter)
# Author-only checklist is deliberately separate from the journal manuscript.
c=newdoc()
c.styles['Normal'].font.size=Pt(10.5);c.styles['Normal'].paragraph_format.space_after=Pt(4);c.styles['Normal'].paragraph_format.line_spacing=1.03
for name in ['Heading 1','Heading 2']:
 c.styles[name].paragraph_format.space_before=Pt(6);c.styles[name].paragraph_format.space_after=Pt(4)
paragraph(c,'Author submission checklist','Title')
paragraph(c,'Journal of Clinical Medicine | Corresponding author: Hui Liu')
paragraph(c,'The manuscript has been assembled with 7 main figures, 14 supplementary figures and 110 references. Submission release requires the author confirmations below and the final combined-upload-size check. This checklist is for the authors and is not part of the scientific manuscript.')
paragraph(c,'Prepared manuscript and supporting materials','Heading 1')
rows=[('Item','Documented status'),('Article structure','JCM Article sections; structured 241-word abstract with Background/Objectives, Methods, Results and Conclusions; 9 keywords; editable Word manuscript and 2 tables.'),('References','110 references in final main-manuscript order. Supplementary citations share that numbering. Journal metadata, software versions and structure-dataset citations are documented in the reference audit.'),('Main and supplementary figures','Main Figures 1–7 and Supplementary Figures S1–S14 are defined. Detailed calibration, sensitivity and complete compound comparisons remain in the supplement.'),('Workflow panels','Four workflow panels redesigned and reassembled: final Figure 1A, Figure 5A, Figure 6A and Figure S12A. Quantitative panels retain their source evidence.'),('Figure 4 repair','Repaired RGB TIFF source independently decoded successfully. Use the repaired Figure 4 export in the final package; the obsolete historical TIFF is excluded.'),('Transferred Results','All 10 transferred Results paragraphs and S11/S12 legends are retained in Additional file 4. Its Davis, State and scBaseCount references use final numbers 73, 75 and 76.'),('Required backmatter','Author contributions, funding, IRB/consent statements, data availability, substantive AI disclosure, acknowledgments and conflicts sections are present. Author-controlled declarations require confirmation below.'),('Upload size','JCM permits no more than 120 MB for all submission files combined. Final aggregate file-size validation remains pending packaging; confirm the final size manifest before upload.')]
table(c,rows)
paragraph(c,'Completed computation and scientific scope','Heading 1')
paragraph(c,'Live BoltzMol 1.0 screening completed for both receptors under the authorized US$0.90 total estimate (US$0.45 per target). Sixteen compounds were submitted per receptor; 12 were scored and 4 were filtered per target. Model filtering is recorded as unscored, not as evidence of absent binding. This cost statement records the authorization/estimate and is not a reconciled billing receipt.')
paragraph(c,'AutoDock Vina completed 96 primary library searches, 3 apo-PDGFRA sensitivity searches and 1 independent imatinib control. Inductive Bio supplied molecular-property predictions. Proto successfully retrieved the PDGFRA and FLT4 AlphaFold structures; its last workspace response still showed Modal disconnected and no deployed tools, so no Proto remote docking or dynamics execution is claimed.')
paragraph(c,'Production ligand–protein molecular dynamics was not completed. The 0.44-ps ligand-free FLT4 OpenMM feasibility calculation does not establish complex stability, convergence or binding free energy and is not presented as production MD evidence. The manuscript’s computational conclusions remain screening and structural hypotheses; protective regulation requires biological testing.')
c.add_page_break()
paragraph(c,'Author confirmations before submission','Heading 1')
items=[
('1. Final version and publication status','Confirm approval by all eight authors of the exact submitted manuscript and supplementary materials, and confirm exclusivity/nonpublication. The clean cover letter does not assert these unverified facts. After confirmation, insert the following declarations into the cover letter:'),
]
for h,t in items:paragraph(c,h,'Heading 2');paragraph(c,t)
paragraph(c,'Neither this manuscript nor any portion is published in another journal or being considered for publication elsewhere.')
paragraph(c,'All authors have approved the manuscript and agree with its submission to JCM.')
paragraph(c,'The first declaration above is equivalent wording for the journal’s requested exclusivity statement. Acknowledge any actual prior MDPI submission and provide its manuscript ID in the portal when available. No negative submission history is presumed.')
for h,t in [
('2. Coauthor email addresses','Supply submission emails for Shaoxin Huang, Yong Yang, Xinyu Zhou, Xinyi Chen, Shan Li, Wenyan Fan and Jianjun Xiong. Hui Liu’s supplied email is LIUHUIOKCA@163.com. Obtain consent for displayed email addresses or follow the journal’s proof-stage nondisplay procedure. Verify current affiliations. ORCID profiles and biographies are encouraged, not required by this checklist.'),
('3. Funding and sponsor roles','Confirm the supplied NSFC grants 82260633 and 82460633 and Jiangxi grant 20262BAC240201. State the sponsors’ actual role in design, execution, analysis/interpretation, writing and publication decisions; add a no-role declaration only if confirmed. Identify publication-cost funding if applicable.'),
('4. Competing interests','Confirm each author’s relevant relationships. Shaoxin Huang and Yong Yang list Jiangxi Jiujieli Life Technology Co., Ltd.; disclose applicable employment, ownership or other interests without inferring them from affiliation alone. Reconcile the final statement with the authors’ actual disclosures.'),
('5. Public-data ethics and consent','Confirm the institutionally applicable secondary-use ethics/consent basis for deidentified public data and any existing exemption determination or relevant local/national basis required by the journal. Do not invent an approval number, waiver or new participant-consent declaration. The manuscript reports no new enrollment, specimens or animal procedures.'),
('6. Final submission checks','Confirm that every named supplementary file is present, all public-data links and any separate large archive are accessible, the final combined portal upload is at most 120 MB, and the exact submitted version matches the approved materials. Add the confirmed cover-letter attestations before upload. Reviewer suggestions belong in the portal, not the cover letter.')]:
 paragraph(c,h,'Heading 2');paragraph(c,t)
paragraph(c,'Official submission guidance','Heading 1')
link(c,'JCM instructions for authors','https://www.mdpi.com/journal/jcm/instructions')
link(c,'JCM aims and scope','https://www.mdpi.com/journal/jcm/about')
link(c,'MDPI research and publication ethics','https://www.mdpi.com/ethics')
link(c,'MDPI ACS reference guide','https://mdpi-res.com/data/mdpi-acs-references-guide-v11-2025.12.pdf')
paragraph(c,'Journal instructions checked on 2 October 2026. The final author attestations and aggregate package-size check are distinct from analytical verification.')
checklist=OUT/'Author_Information/Author_Submission_Checklist.docx';c.save(checklist)
meta={'cover_letter':str(letter),'checklist':str(checklist),'cover_letter_missing_only_author_attestations':True,'required_attestations_not_invented':True,'production_MD_status':'not_completed','screening_status':'both_succeeded','cost_scope':'authorized_estimate_USD_0.90_not_reconciled_receipt','portal_size_status':'pending_final_packaging','references':110,'main_figures':7,'supplementary_figures':14,'workflow_panels':4}
(B/'Author_Documents_Build_Report.json').write_text(json.dumps(meta,indent=2));print(json.dumps(meta,indent=2))
