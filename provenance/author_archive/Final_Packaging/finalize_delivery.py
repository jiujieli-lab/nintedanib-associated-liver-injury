from pathlib import Path
import shutil,hashlib,json,zipfile,csv,re
from docx import Document
import fitz
R=Path(__file__).resolve().parents[1];S=R/'submission';D=R.parent.parent/'JCM_Delivery';D.mkdir(exist_ok=True)
shutil.copy2(R/'manuscript/render_final/JCM_Manuscript.pdf',S/'Manuscript/JCM_Manuscript.pdf')
# Reading copies are distinct from the files cited as numbered supplements.
read=S/'Author_Information/Reading_Copies';read.mkdir(exist_ok=True)
p=R/'journal_audit/targeted_supplement_render/Additional_file_4_Targeted_Analyses.pdf'
if p.is_file():shutil.copy2(p,read/p.name)
for fn in ['Final_Figure_Asset_Map.json','manuscript/Assembly_Report.json','manuscript/Final_Reference_Number_Map.json']:
 p=R/fn;shutil.copy2(p,S/'Author_Information'/p.name)
# All journal-facing files except author-only notes and duplicate PDF reading copies.
def portal_files():
 return [p for p in S.rglob('*') if p.is_file() and 'Author_Information' not in p.parts and p.name!='JCM_Manuscript.pdf' and p.name!='Additional_file_7_Screening_Methods_and_Results.pdf' and p.name not in ('README.txt','SHA256SUMS.txt')]
files=portal_files();total=sum(p.stat().st_size for p in files);assert total<120_000_000,total
manifest={'journal':'Journal of Clinical Medicine','guide_checked':'2026-10-02','portal_limit_bytes':120_000_000,'selected_upload_bytes':total,'selected_upload_MB':round(total/1e6,3),'selected_files':[{'file':str(p.relative_to(S)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(files)],'author_attestations_required':True,'production_MD_completed':False,'paid_screening':'Both jobs succeeded under authorized USD0.90 estimate; final billing not reconciled','docking_searches':{'primary':96,'apo_sensitivity':3,'separate_redocking_control':1}}
(S/'Author_Information/Portal_Upload_Manifest.json').write_text(json.dumps(manifest,indent=2))
# Resolve the checklist's packaging-only pending item; leave human attestations unchanged.
p=S/'Author_Information/Author_Submission_Checklist.docx';doc=Document(p)
for q in doc.paragraphs:
 q.text=q.text.replace('and the final combined-upload-size check','listed in this checklist').replace('below listed in this checklist','listed in this checklist').replace('The final author attestations and aggregate package-size check are distinct from analytical verification.','The author attestations remain distinct from the completed analytical and package-size checks.')
for table in doc.tables:
 for row in table.rows:
  if row.cells[0].text=='Upload size':row.cells[1].text=f'The selected journal-upload set totals {total/1e6:.2f} MB, below the 120 MB aggregate limit. Portal_Upload_Manifest.json identifies the files. Author-only material, duplicate reading PDFs and separate artwork/reproducibility archives are excluded from this set.'
doc.save(p)
(S/'README.txt').write_text(f'''JOURNAL OF CLINICAL MEDICINE SUBMISSION MATERIALS

Manuscript/JCM_Manuscript.docx is the editable article with7 main figures,2 tables and110 references. JCM_Manuscript.pdf is a reading copy. Main_Figures contains the7 separately supplied verifiedTIFF files. Supplementary_Materials contains numbered Additionalfiles1–9. Additionalfile2 contains14 supplementary figures; Additionalfiles7–9 contain the executed compound-screening and docking analyses.

The selected journal-upload set is {total/1e6:.2f} MB against the120MB aggregate guide limit. Use Author_Information/Portal_Upload_Manifest.json to identify that set. Do not add the separate high-resolution artwork archive, author reproducibility archive, author checklist or duplicate reading PDFs to this portal selection without recalculating its size.

Before submission, complete Author_Information/Author_Submission_Checklist.docx. Required author-controlled items include exact-version approval, publication/submission status, coauthor contact details, sponsor roles and applicable competing-interest/ethics declarations. Insert confirmed cover-letter attestations; none have been invented. The cover letter is a clean scientific letter awaiting these factual confirmations.

Production target–ligand molecular dynamics has NOT been completed. The article reports actual AI screening, molecular-property predictions, rigid docking, structural sensitivity and pre-existing public-data analyses. It contains no invented MD trajectory or binding-free-energy result. The author archive contains the feasibility findings and prospective simulation protocol.

Figure4.tiff is the repaired RGB1200dpi export. New screening figures are600dpi. Original raster figures retain their native pixels and honest resolution; changing a resolution tag would not add source detail. All numerical panels, nonsignificant results and source-record dispositions are retained in the relevant main or supplementary materials.

Source identifiers in Additionalfile3 are mapped to final figure/panel numbers on page93 of Additionalfile1 and in the complete crosswalk in Additionalfile9. Source filenames and numerical values were preserved.
''')
# Final checksums and archive integrity.
def zipdir(src,out):
 hashes=[]
 for p in sorted(src.rglob('*')):
  if p.is_file() and p.name!='SHA256SUMS.txt':hashes.append(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+str(p.relative_to(src)))
 (src/'SHA256SUMS.txt').write_text('\n'.join(hashes)+'\n')
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
  for p in sorted(src.rglob('*')):
   if p.is_file():z.write(p,p.relative_to(src))
 with zipfile.ZipFile(out) as z:assert z.testzip() is None
 return {'file':out.name,'bytes':out.stat().st_size,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'files':len(hashes)+1,'zip_crc_pass':True}
# Scientific file count and readability gates.
assert len(list((S/'Main_Figures').glob('Figure_*.tiff')))==7
assert len(fitz.open(S/'Supplementary_Materials/Additional_file_2_Supplementary_Figures.pdf'))==14
assert len(fitz.open(S/'Supplementary_Materials/Additional_file_1_Supplementary_Information.pdf'))==93
assert len(Document(S/'Manuscript/JCM_Manuscript.docx').tables)==2
refs=json.loads((R/'manuscript/Final_References.json').read_text());assert len(refs)==110 and [x['number'] for x in refs]==list(range(1,111))
for p in (S/'Supplementary_Materials').glob('*.zip'):
 with zipfile.ZipFile(p) as z:assert z.testzip() is None,p
# Copies for immediate access use English names and exact validated bytes.
for src,name in [(S/'Manuscript/JCM_Manuscript.docx','JCM_Manuscript.docx'),(S/'Manuscript/JCM_Manuscript.pdf','JCM_Manuscript.pdf'),(S/'Main_Figures/Figure_4.tiff','Figure_4_Repaired_1200dpi.tiff')]:shutil.copy2(src,D/name)
report={'portal_upload':manifest,'gates':{'main_figures':7,'supplementary_figures':14,'extended_supplement_pages':93,'main_references':110,'production_md_claimed':False},'archive_outputs':[]}
# Author archive must already be staged. Copy current release-specific audit files now.
for p in [R/'Final_Figure_Asset_Map.json',S/'Author_Information/Portal_Upload_Manifest.json',R/'packaging/Publication_Text_QA.json']:
 q=R/'author_archive/Final_Packaging'/p.name;q.parent.mkdir(exist_ok=True,parents=True);shutil.copy2(p,q)
for src,name in [(S,'JCM_Submission_Package.zip'),(R/'artwork_archive','JCM_High_Resolution_Artwork.zip'),(R/'author_archive','JCM_Author_Reproducibility_Archive.zip')]:report['archive_outputs'].append(zipdir(src,D/name))
(R/'packaging/Final_Delivery_QA.json').write_text(json.dumps(report,indent=2));(D/'Delivery_Manifest.json').write_text(json.dumps(report,indent=2));print(json.dumps({'portal_MB':round(total/1e6,3),'archives':report['archive_outputs']},indent=2))
