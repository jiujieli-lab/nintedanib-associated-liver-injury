from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'deliverables'

def base(title=None):
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.85);s.bottom_margin=Inches(.85);s.left_margin=s.right_margin=Inches(1)
 for style in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[style];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(12)
 d.styles['Normal'].paragraph_format.line_spacing=1.15;d.styles['Normal'].paragraph_format.space_after=Pt(7)
 for st in ['Heading 1','Heading 2']:
  d.styles[st].paragraph_format.keep_with_next=True;d.styles[st].paragraph_format.space_before=Pt(12);d.styles[st].font.bold=True
 d.styles['Title'].font.size=Pt(14);d.styles['Title'].font.bold=True
 for el in list(d.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 for el in d.styles.element.iter(qn('w:rFonts')):
  for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme']:
   el.attrib.pop(qn('w:'+attr),None)
  el.set(qn('w:ascii'),'Times New Roman');el.set(qn('w:hAnsi'),'Times New Roman')
 foot=s.footer.paragraphs[0];foot.alignment=2;r=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.append(fld)
 d.core_properties.author='Shaoxin Huang';d.core_properties.last_modified_by='Shaoxin Huang'
 if title:d.core_properties.title=title;d.add_paragraph(title,'Title')
 return d

# Source-aligned cover letter; all substantive submission declarations preserved.
letter=(ROOT/'work/editorial/cover_letter_new.txt').read_text()
d=base();d.core_properties.title='Cover letter for nintedanib target prioritization manuscript'
for block in re.split(r'\n\s*\n',letter.strip()):
 p=d.add_paragraph(block)
 if block.startswith('Dear Editors'):p.paragraph_format.space_before=Pt(5)
d.save(OUT/'Manuscript/Cover_Letter.docx')

# Supplementary prose includes completed work only, with independent references.
d=base('Targeted biochemical and virtual intervention analyses')
d.add_paragraph('Clinical proteomics-guided network integration prioritizes PDGFRA and FLT4 for investigation of nintedanib-associated liver injury')
d.add_paragraph('This supplement reports the completed target-focused reanalysis and model-conditional network interventions. Figure 4 provides the endpoint-resolved biochemical evidence; Supplementary Figure S11 provides all targeted network-intervention results. Numerical tables N1–N17 are supplied in Additional file 5, and complete gene-level and matched-null outputs N18–N19 accompany Additional file 6. References are numbered independently within this supplement.')
# Binding sections.
md=(ROOT/'work/binding_analysis/Manuscript_Binding_Inserts.md').read_text()
parts=re.split(r'^## ',md,flags=re.M)
for chunk in parts[1:]:
 title,*rest=chunk.split('\n');body='\n'.join(rest).strip()
 if title.startswith(('Proposed Discussion','Required disclosure','Figure 4 legend')):continue
 title=title.replace('Methods: ','').replace('Results: ','').replace('Supplementary Methods: ','').replace('Supplementary Results: ','')
 title=title[0].upper()+title[1:]
 title={'Experimental target engagement is supported, whereas relative selectivity depends on the assay':'Biochemical engagement and assay dependent selectivity','Supplementary connector-based physicochemical predictions':'Molecular property prediction methods','Supplementary physicochemical context for prospective simulations':'Molecular property prediction results'}.get(title,title)
 d.add_paragraph(title,'Heading 1')
 for p in re.split(r'\n\s*\n',body):
  if p:d.add_paragraph(p)
d.add_paragraph('The direct-binding records originate from Davis et al. [1]. Activities 7574898 and 7573179 in CHEMBL1908390 report PDGFRA Kd=16 nM and FLT4 Kd=95 nM, respectively. The accession-linked source document and individual activity records are retained with the analysis.')
# Virtual methods and complete results.
v=(ROOT/'work/new_virtual/manuscript_insertion.txt').read_text()
m=v.split('PROPOSED METHODS INSERTION (after donor-balanced network methods)\n',1)[1].split('\nPROPOSED CONCISE MAIN RESULTS INSERTION',1)[0].strip()
lines=m.split('\n');d.add_paragraph('Single and dual target network intervention methods','Heading 1')
for p in lines[1:]:
 if p.strip():
  p=p.replace('[95]','[2]')
  # Express the estimand in ordinary prose without pseudo-LaTeX subscripts.
  p=p.replace('With baseline and intervened propagation vectors f0 and fS, the primary magnitude was M(S)=sum_g w_g |fS_g−f0_g| / sum_g w_g f0_g over the 13 clinical genes.', 'For each of the 13 clinical genes, the absolute difference between intervened and baseline propagation scores was multiplied by its clinical weight. The sum of these products was divided by the clinical-weighted baseline score to obtain the primary relative magnitude.')
  p=p.replace('The exact joint residual was I=fAB−fA−fB+f0;', 'The exact joint residual was the dual-target score minus each single-target score plus the baseline score;')
  d.add_paragraph(p)
d.add_paragraph('Matrix provenance retained GEO accession GSE115469, the deposited normalized measurement convention, original cell identifiers and the five donor labels. The existing data were not pooled with differently quantified counts or relabelled as an independently processed scBaseCount dataset. This provenance design follows the comparability and metadata considerations described by Youngblut et al. [3].')
d.add_paragraph('Target pair intervention results','Heading 1')
r=v.split('PROPOSED CONCISE MAIN RESULTS INSERTION (after integrated prioritization)\n',1)[1].split('\nOPTIONAL VERY SHORT ABSTRACT ADDITION',1)[0].strip().split('\n',1)[1]
d.add_paragraph(r)
qc=json.loads((ROOT/'work/new_virtual/validation_summary.json').read_text())
d.add_paragraph('Reconstruction reproduced all 50 archived primary target–compartment propagation scores, with maximum absolute difference 4.07×10−20. The largest linear-solve residual was 9.71×10−17; sham deletion and graph restoration both returned zero displacement. The latter confirms numerical reversibility of the operator and is not a biological rescue experiment. All full-data, comparator, sham and donor-omission calculations are retained in the 111-row intervention table. The fixed attenuation series scaled shared incident edges once, avoiding double attenuation of the edge between the two targets.')
d.add_paragraph('The macrophage joint-residual comparison had nominal P=0.03370 and BH q=0.10110; the corresponding hepatocyte and endothelial q values were 0.99138 and 0.24974. All nine primary magnitude tests had q≥0.52941. Complete target and compartment results are provided in the accompanying tables.')
d.add_paragraph('Supplementary Figure S11 legend','Heading 1')
legend=v.split('FIGURE LEGEND\n',1)[1].split('\nSUPPLEMENTARY TABLE / SOURCE DATA CONTENTS',1)[0].strip()
for p in legend.split('\n'):
 if p.strip():d.add_paragraph(p)
d.add_paragraph('Supplementary table and source data guide','Heading 1')
manifest=json.loads((ROOT/'work/tables/table_manifest.json').read_text())
for x in manifest:
 name=x['sheet'].split('_',1)[1].replace('_',' ') if x.get('sheet') else x['filename']
 d.add_paragraph(f"Table {x['number']}. {name}. {x['rows']:,} records. "+x.get('note','')+(' The complete compressed CSV is supplied in Additional file 6.' if not x.get('sheet') else ''))
d.add_paragraph('Execution and reproducibility','Heading 1')
d.add_paragraph('The standalone code archive includes the three original normalized compartment matrices, donor metadata, clinical weights, archived primary graph scores, pharmacology source rows and raw Inductive Bio responses required to reproduce the additions. Input and output SHA-256 manifests preserve the analysed version. The calculation and figure scripts document the model specifications, controls and numerical checks for each completed analysis.')
d.add_paragraph('References','Heading 1')
refs=[
'2. Adduri AK, Gautam D, Bevilacqua B, Naghipourfar M, Imran A, Shah R, et al. Predicting cellular responses to perturbation across diverse contexts with State. Cell. 2026;189:5914-5931. https://doi.org/10.1016/j.cell.2026.07.052.',
'3. Youngblut ND, Carpenter C, Nayebnazar A, Adduri A, Shah R, Ricci-Tam C, et al. scBaseCount: An AI agent-curated, standardized, auto-updated single-cell data repository. Cell. 2026;189:5932-5944. https://doi.org/10.1016/j.cell.2026.08.025.',
'1. '+(ROOT/'work/binding_analysis/Kd_Primary_Reference.txt').read_text().strip()]
for x in sorted(refs,key=lambda x:int(x.split('.',1)[0])):
 p=d.add_paragraph(x);p.paragraph_format.left_indent=Inches(.3);p.paragraph_format.first_line_indent=Inches(-.3)
for p in d.paragraphs:
 if re.search(r'10[−-]\d+',p.text):
  txt=p.text
  for r in list(p.runs):p._p.remove(r._r)
  offset=0
  for m in re.finditer(r'10([−-]\d+)',txt):
   p.add_run(txt[offset:m.start(1)]);p.add_run(m[1]).font.superscript=True;offset=m.end()
  p.add_run(txt[offset:])
d.save(OUT/'Supplementary_Materials/Additional_file_4_Targeted_Analyses.docx')
print('Created Cover_Letter.docx and Additional_file_4_Targeted_Analyses.docx')
