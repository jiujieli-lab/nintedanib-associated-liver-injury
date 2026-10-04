from pathlib import Path
from copy import deepcopy
from io import BytesIO
from datetime import datetime,timezone
import re,json,difflib,hashlib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/'work/source/转化医学投稿20260922/JTM_Manuscript20260924.docx'
OUT=ROOT/'deliverables/Manuscript';OUT.mkdir(parents=True,exist_ok=True)
doc=Document(SRC)
original_nodes={el:''.join(el.itertext()) for el in []}
def text_of(el):return ''.join(el.iter(qn('w:t'))[0].itertext()) if False else ''.join(n.text or '' for n in el.iter(qn('w:t')))
original_nodes={el:text_of(el) for el in doc.element.body.iter(qn('w:p'))}
original_math={}
for el in original_nodes:
 chars=[]
 for r in el.iter(qn('w:r')):
  va=r.find('w:rPr/w:vertAlign',r.nsmap)
  flag=va.get(qn('w:val')) if va is not None else None
  chars.extend([flag]*len(text_of(r)))
 original_math[el]=chars
def restore_math(p_el,flags):
 offset=0
 for r in list(p_el.iter(qn('w:r'))):
  if any(a.tag==qn('w:del') for a in r.iterancestors()):continue
  ts=r.findall(qn('w:t'))
  if len(ts)!=1:continue
  txt=ts[0].text or '';local=flags[offset:offset+len(txt)];offset+=len(txt)
  if not any(local):continue
  groups=[]
  for ch,flag in zip(txt,local):
   if not groups or groups[-1][0]!=flag:groups.append([flag,ch])
   else:groups[-1][1]+=ch
  parent=r.getparent();idx=parent.index(r)
  for flag,part in groups:
   nr=deepcopy(r);nr.find(qn('w:t')).text=part;nr.find(qn('w:t')).set(qn('xml:space'),'preserve')
   if flag:
    pr=nr.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');nr.insert(0,pr)
    for va in list(pr.findall(qn('w:vertAlign'))):pr.remove(va)
    va=OxmlElement('w:vertAlign');va.set(qn('w:val'),flag);pr.append(va)
   parent.insert(idx,nr);idx+=1
  parent.remove(r)
orig=list(doc.paragraphs);changes=[]
def replace(p,text):
 old=p.text
 if old==text:return
 changes.append({'before':old,'after':text})
 rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for ch in list(p._p):
  if ch.tag!=qn('w:pPr'):p._p.remove(ch)
 r=p.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
def add_after(anchor,text,style='Normal'):
 el=OxmlElement('w:p');anchor._p.addnext(el);p=Paragraph(el,anchor._parent);p.style=style;p.add_run(text);return p
def insert_block(anchor,items):
 for style,text in items:anchor=add_after(anchor,text,style)
 return anchor
# Stable remap of ORIGINAL main figures only. Supplement identifiers are untouched.
figmap={1:1,2:2,3:3,4:5,5:6,6:7}
for p in doc.paragraphs:
 t=re.sub(r'\b(Figure|Fig\.) ([1-6])(?=[A-Z\s,;.\)])',lambda m:m[1]+' '+str(figmap[int(m[2])]),p.text)
 if t!=p.text:replace(p,t)
# Introductory context and explicit schematic callouts.
replace(orig[22],orig[22].text.replace('The analysis followed','The study sequence is summarized in Fig. 1A. The analysis followed'))
replace(orig[37],orig[37].text+' Figure 5A summarizes the localization and numerical-calibration sequence.')
replace(orig[42],orig[42].text+' Figure 6A summarizes the clinically weighted network integration.')
replace(orig[52],orig[52].text+' The independent evaluation design is shown in Fig. 7A.')
replace(orig[23],orig[23].text.replace('Figures S1–S10 are retained in their original order in Additional file 2.','Supplementary Figures S1–S11 are provided in Additional file 2. Target-focused binding and virtual-intervention methods and results are detailed in Additional file 4, with their complete numerical tables in Additional file 5.'))
# New methods describe only completed analyses.
binding_methods=[('Heading 2','Target-focused biochemical binding evidence and molecular-property predictions'),('Normal',
'We further resolved the PDGFRA and FLT4 records within the frozen ChEMBL extraction by endpoint, assay and source document. Direct dissociation constants (Kd), inhibition constants (Ki) and half-maximal inhibitory concentrations (IC50) were analysed separately. The original eligibility rules remained primary. A labelled assay-format sensitivity additionally admitted exact, validity-clean human single-protein records whose descriptions specified recombinant kinase assays but whose BioAssay Ontology format was recorded generically. Duplicate-flagged, nonquantitative kinetic and single-concentration inhibition records were retained in the disposition table and excluded from potency summaries. We reported every eligible concentration, observed range, document-balanced median and leave-one-document-out median; matched-document contrasts compared the same endpoint without assuming matched assay conditions. No inferential pooling across assay types or conventional molecular docking was performed. Direct Kd observations were traced to the original kinase-binding study [97].'),('Normal',
'The connected Inductive Bio service predicted nintedanib LogD at pH 7.4 and acidic/basic pKa from the stereochemically specified neutral-parent ChEMBL SMILES. Resolved model versions were mcp_public_logd 1.7.0 and mcp_public_apka/mcp_public_bpka 1.4.0. The exact input, returned values, model configuration identifiers and applicability flags were archived. Provider-returned lower and upper bounds were retained without assigning an unreported confidence level. These physicochemical predictions describe the ligand and were not used as target-binding or molecular-dynamics measurements (Additional files 4 and 5).')]
insert_block(orig[35],binding_methods)
virtual_methods=[('Heading 2','Single and joint target interventions in the clinical network'),('Normal',
'The extension adopted perturbation-delta benchmarking and cell-context applicability principles from State [95] and explicit matrix and donor provenance from scBaseCount [96]. We retained the deposited normalized GSE115469 matrices and their five donor identifiers; no State checkpoint, scBaseCount-derived observations or new experimental perturbation data entered the analysis. Within the fixed hepatocyte, endothelial and macrophage gene universes, we reconstructed the original equal-donor Ledoit–Wolf positive top-20-neighbour graph with restart probability 0.5 and discovery DO-versus-HV clinical-protein weights. All incident edges were removed for PDGFRA, FLT4 or both targets, isolated nodes received self-loops, and the column-stochastic transition was rebuilt before propagation. The primary magnitude was the clinical-weighted absolute difference between intervened and baseline propagation scores divided by the weighted baseline score. The joint residual was the dual-intervention score minus both single-intervention scores plus the baseline. KDR and FGFR1 supplied comparator deletions; sham deletion and reapplication of the original graph checked numerical consistency. These operations change network proximity and do not estimate transcript abundance or biological rescue.'),('Normal',
'Each target was matched to 50 background genes on standardized log1p equal-donor expression, logit detection with a 0.005 pseudocount and log1p weighted degree, excluding the 13 clinical genes and 17 pharmacological candidates. All single-target controls and unique unordered matched pairs were evaluated: 1275, 1939 and 1275 pairs in hepatocyte, endothelial and macrophage networks, respectively. Upper-tail empirical probabilities used a plus-one correction. BH adjustment covered nine intervention-by-compartment magnitude tests and, separately, three joint-residual tests. Each donor was omitted in turn and the graph rebuilt within the inherited feature universe. A fixed 0%, 25%, 50%, 75% and 100% attenuation series scaled every incident edge once. Donor-omission ranges describe within-atlas sensitivity. Settings were fixed before calculation on 2 October 2026, after the original candidate ranking was known, and were not prospectively preregistered. Supplementary Fig. S11 and Additional files 4 and 5 retain all outcomes.')]
insert_block(orig[50],virtual_methods)
# New results.
binding_results=[('Heading 2','Experimental biochemical records support engagement of both prioritized receptors'),('Normal',
'The frozen pharmacology set contained 11 quantitative records for the two priorities: seven for PDGFRA and four for FLT4. Direct binding measurements reported Kd values of 16 nM for PDGFRA and 95 nM for FLT4 in the same kinase-binding study [97], complementing the cellular-context ranking with measured receptor-binding evidence (Fig. 4A,B). PDGFRA IC50 values spanned 1.9–560 nM across five source documents, with a median of 18 nM; FLT4 IC50 values spanned 5–34 nM across three documents, with a median of 13 nM. A separate PDGFRA Ki was 8.4 nM. Thus, both receptors had submicromolar experimental pharmacological support, while direct affinity and inhibitory potency remained distinct endpoints.'),('Normal',
'Within the two primary documents reporting both receptors, PDGFRA-to-FLT4 IC50 ratios were 3.60 and 4.54. Three additional records entered the labelled assay-format sensitivity; the expanded IC50 medians were 11 nM for PDGFRA and 5 nM for FLT4, and two additional matched-document ratios were 0.63 and 0.72 (Fig. 4C,D). The variable relative ordering and document-omission profiles support engagement of both receptors without establishing uniform selectivity for either receptor (Fig. 4E,F). These endpoint-resolved findings elaborate the original pharmacology source rather than add an independent validation cohort. Inductive Bio returned LogD7.4 2.523, acidic pKa 10.945 and basic pKa 7.712, with no low-confidence or out-of-domain flags; the full returned bounds and configurations are reported in Additional files 4 and 5.')]
insert_block(orig[77],binding_results)
virtual_results=[('Heading 2','Target-pair perturbations characterize clinical-network sensitivity'),('Normal',
'Single and joint PDGFRA/FLT4 edge deletions generated numerically consistent outputs across donor omissions. The joint deletion changed the clinical-weighted propagation score by 0.000336%, 0.02154% and 0.03848% in hepatocyte, endothelial and macrophage networks, respectively; none of the nine primary effects exceeded matched-control backgrounds after correction (all q≥0.529; Supplementary Fig. S11A–F). Joint residuals also remained nonsignificant after three-test correction (q=0.991, 0.250 and 0.101; Supplementary Fig. S11G,H). Reciprocal receptor-score redistribution was descriptive (Supplementary Fig. S11I). PDGFRA was detected in 8, 33 and 2 cells in the respective compartments, compared with 4, 150 and 13 for FLT4. The extension therefore defines model-level sensitivity and cellular-coverage requirements while retaining the original interpretation of PDGFRA and FLT4 as priorities for receptor-specific testing.')]
insert_block(orig[89],virtual_results)
# Interpret the added evidence once, alongside the original compartment-specific strategy.
insert_block(orig[109],[('Normal',
'The endpoint-resolved pharmacology adds a direct biochemical basis for this experimental choice: both receptors have measured nanomolar binding affinity, alongside inhibitory activity across multiple source documents. Assay-dependent relative potency argues for measuring occupancy or proximal receptor phosphorylation at the same intracellular nintedanib exposure in the selected hepatic systems. The virtual-intervention extension then asks a different question—how changing each receptor node alters the clinically weighted network. Its calibrated results support neither a synergistic interaction nor a directional PDGFRA–FLT4 regulatory relationship. State emphasizes evaluation against held-out measured perturbations within the model’s supported intervention domain [95], while scBaseCount provides principles for improving reference-cell provenance and coverage [96]. Together, these considerations focus the next validation on target-expressing hepatic contexts and experimentally measured perturbation responses.')])
# Abstract remains bounded by the observed evidence and <=350 words.
replace(orig[9], 'Methods: We integrated active-comparator pharmacovigilance, 13 human serum injury proteins, nintedanib pharmacology and cellular responses, and a five-donor liver atlas. Numerical auditing, donor-balanced shrinkage networks, matched-null propagation, six-algorithm comparison, donor omission and weight perturbation evaluated candidate robustness. Endpoint-resolved biochemical records and single/joint target-network interventions characterized the two priorities. Frozen rankings were challenged in independent liver profiles from 22 mice; pulmonary datasets defined antifibrotic readouts.')
replace(orig[10], 'Results: PDGFRA and FLT4 ranked first and second under four aggregation rules and remained among the top five in all 5000 weight perturbations. PDGFRA remained first after each donor omission; FLT4 ranked second to fifth. Reported direct Kd values were 16 and 95 nM, respectively, supporting biochemical engagement. The narrow hepatic reporting odds ratio for nintedanib versus pirfenidone was 3.997 (95% confidence interval 3.226–4.951), with temporal heterogeneity. FBP1 distinguished drug-related from other acute liver injury within the released serum panel (Cliff’s delta −0.505; q=0.000469). Numerical auditing identified noise-scale displacement in 132/210 virtual-knockout runs, correcting the earlier FGFR1 prioritization. Neither competitive ranking calibration nor targeted network interventions established significant enrichment, and rankings were not significantly associated with external hepatic expression effects. Independent liver analysis identified decreased Kdr expression within the candidate family (log2 difference −0.280; exact-test q=0.0386; genome-wide q=0.119). Pulmonary responses varied by model; calibrated global reversal did not survive multiplicity correction.')
# Preserve original no-new-experiments statement; consolidate actual computational assistance in declarations.
replace(orig[61],'Reproducibility and statistical units')
replace(orig[63],'No new human samples, cell experiments or animal experiments were generated. The supplementary perturbation-and-rescue protocol is prospective and contributes no observed result.')
insert_block(orig[135],[('Heading 2','Use of artificial intelligence'),('Normal',
'OpenAI ChatGPT with Codex tools assisted literature retrieval, analytical code development and review, quantitative figure preparation, manuscript organization, language revision and consistency checking during September–October 2026. Inductive Bio molecular-property models and their versions are documented in Methods. Computational outputs were evaluated against their source data and numerical checks and were not treated as experimental observations or independent validation. The authors remain responsible for data provenance, analytical choices, interpretation and the final manuscript.')])
replace(orig[126], 'The public datasets analysed in this study are accessible through the repositories and source publications identified in Tables 1 and 2. Extended methods and the prospective biological protocol are provided in Additional file 1; Supplementary Figures S1–S11 in Additional file 2; the original complete supplementary tables in Additional file 3; and the targeted biochemical and virtual-intervention methods, results and tables in Additional files 4 and 5. Additional file 6 contains the new analysis scripts, input specifications, source-linked outputs and numerical checks. The accompanying source reproducibility archive retains the original five data/code packages and figure-composition records. This documentation distinguishes archived source results, reconstructed auxiliary files and calibrated analyses. The raw Inductive Bio responses include exact inputs and resolved model versions.')
replace(orig[138], 'Additional file 2. Supplementary Figures S1–S11. Additional_file_2_Supplementary_Figures.pdf. All ten original supplementary figures and the targeted single/joint network-intervention figure, retaining every original data panel.')
insert_block(orig[139],[('Normal','Additional file 4. Targeted analyses. Additional_file_4_Targeted_Analyses.docx. Extended biochemical-binding, molecular-property and virtual-intervention methods, results and the complete Figure S11 legend.'),('Normal','Additional file 5. Targeted analysis tables. Additional_file_5_Targeted_Analysis_Tables.xlsx. Complete endpoint-resolved biochemical records and virtual-intervention effect estimates, null calibration and sensitivity results.'),('Normal','Additional file 6. Targeted analysis code and source data. Additional_file_6_Targeted_Analysis_Code.zip. Executable analysis and figure scripts, frozen specifications, full gene-level outputs, matched-null results and input/output checksums.')])
replace(orig[137], orig[137].text+' References in this supplement are numbered independently.')
# New figure legend placed by final figure number.
binding_legend='(A,B) Individual reported PDGFRA and FLT4 assay values. Circles, squares and triangles denote IC50, Kd and Ki, respectively. Filled symbols identify the frozen primary set; open symbols identify records admitted only by the labelled assay-format sensitivity. (C) Within-document IC50 ratios; each document supplies one descriptive contrast, without assumed equivalence of assay conditions. (D) Endpoint-specific median pChEMBL and observed range; segments are not confidence intervals. (E) Disposition of all 26 target records, including excluded records. (F) Median pIC50 after each source-document omission; black ticks show the complete-set medians. Concentrations are in nM. The primary set contains seven PDGFRA and four FLT4 records; the sensitivity adds one and two records, respectively. Kd and Ki remain separate from IC50. These are reanalysed experimental records, not newly performed binding assays, docking scores or molecular-dynamics results. Full identifiers and eligibility reasons are supplied in Additional file 5.'
insert_block(orig[146],[('Heading 2','Figure 4 Endpoint-resolved biochemical evidence supports engagement of PDGFRA and FLT4'),('Normal',binding_legend)])
# Append verified source references before renumbering by first appearance.
newrefs=[
'95. Adduri AK, Gautam D, Bevilacqua B, Naghipourfar M, Imran A, Shah R, et al. Predicting cellular responses to perturbation across diverse contexts with State. Cell. 2026;189:5914-5931. https://doi.org/10.1016/j.cell.2026.07.052.',
'96. Youngblut ND, Carpenter C, Nayebnazar A, Adduri A, Shah R, Ricci-Tam C, et al. scBaseCount: An AI agent-curated, standardized, auto-updated single-cell data repository. Cell. 2026;189:5932-5944. https://doi.org/10.1016/j.cell.2026.08.025.',
]
ref97=ROOT/'work/binding_analysis/Kd_Primary_Reference.txt'
if not ref97.exists():raise RuntimeError('Verified direct binding reference not yet available')
newrefs.append('97. '+ref97.read_text().strip().removeprefix('97. '))
last=doc.paragraphs[-1]
for r in newrefs:last=add_after(last,r,'Normal')
# Number references in main document by first citation, including table text.
refheading=next(p for p in doc.paragraphs if p.text=='References')
ref_nodes=[];past=False
for el in list(doc.element.body):
 if el is refheading._p:past=True;continue
 if past and el.tag==qn('w:p') and re.match(r'^\d+\. ',text_of(el)):ref_nodes.append(el)
refs={int(re.match(r'^(\d+)\.',text_of(el))[1]):el for el in ref_nodes}
pattern=re.compile(r'\[(\d+(?:\s*(?:,|–|-)\s*\d+)*)\]')
def nums(t):
 out=[]
 for x in t.split(','):
  y=re.split('[–-]',x.strip());out.extend(range(int(y[0]),int(y[1])+1) if len(y)==2 else [int(y[0])])
 return out
order=[]
for el in list(doc.element.body):
 if el is refheading._p:break
 for p in el.iter(qn('w:p')) if el.tag!=qn('w:p') else [el]:
  for m in pattern.finditer(text_of(p)):
   for n in nums(m[1]):
    if n not in refs:raise RuntimeError('unknown reference '+str(n))
    if n not in order:order.append(n)
for n in sorted(refs):
 if n not in order:raise RuntimeError('uncited reference '+str(n))
renum={n:i+1 for i,n in enumerate(order)}
for el in list(doc.element.body):
 if el is refheading._p:break
 for p_el in el.iter(qn('w:p')) if el.tag!=qn('w:p') else [el]:
  p=Paragraph(p_el,doc._body);old=p.text
  new=pattern.sub(lambda m:'['+', '.join(str(renum[n]) for n in nums(m[1]))+']',old)
  if new!=old:replace(p,new)
anchor=refheading._p
for oldnum in order:
 el=refs[oldnum];p=Paragraph(el,doc._body);replace(p,re.sub(r'^\d+\.',str(renum[oldnum])+'.',p.text));anchor.addnext(el);anchor=el
for ref_el in refs.values():
 p=Paragraph(ref_el,doc._body);p.paragraph_format.left_indent=Inches(.3);p.paragraph_format.first_line_indent=Inches(-.3)
# Preserve source format; black headings and native line numbering.
for st in doc.styles:
 if st.type==1 and (st.name.startswith('Heading') or st.name=='Title'):
  st.font.color.rgb=RGBColor(0,0,0);st.font.name='Times New Roman'
for el in doc.styles.element.iter(qn('w:rFonts')):
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme']:el.attrib.pop(qn('w:'+attr),None)
 el.set(qn('w:ascii'),'Times New Roman');el.set(qn('w:hAnsi'),'Times New Roman')
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
# Fit source tables inside the text area with explicit gutters and intact rows.
for ti,table in enumerate(doc.tables):
 widths=[1.70,1.85,2.95] if ti==0 else [1.05,2.10,1.75,1.60]
 table.autofit=False
 for col,w in zip(table.columns,widths):col.width=Inches(w)
 pr=table._tbl.tblPr
 tw=pr.find(qn('w:tblW'))
 if tw is None:tw=OxmlElement('w:tblW');pr.append(tw)
 tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9360')
 borders=pr.find(qn('w:tblBorders'))
 if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
 for side in ['top','left','bottom','right','insideH','insideV']:
  b=borders.find(qn('w:'+side))
  if b is None:b=OxmlElement('w:'+side);borders.append(b)
  b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9')
 for ri,row in enumerate(table.rows):
  trpr=row._tr.get_or_add_trPr()
  if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
  if ri==0 and trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
  for ci,(cell,w) in enumerate(zip(row.cells,widths)):
   cell.width=Inches(w);tcpr=cell._tc.get_or_add_tcPr();mar=tcpr.find(qn('w:tcMar'))
   if mar is None:mar=OxmlElement('w:tcMar');tcpr.append(mar)
   for side,val in [('top',60),('bottom',60),('left',95),('right',95)]:
    m=mar.find(qn('w:'+side))
    if m is None:m=OxmlElement('w:'+side);mar.append(m)
    m.set(qn('w:w'),str(val));m.set(qn('w:type'),'dxa')
   for p in cell.paragraphs:
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    for r in p.runs:r.font.size=Pt(10.5);r.font.name='Times New Roman'
# Figure legends must fit the journal limit and abstract <=350 words.
abstract=' '.join(orig[i].text for i in range(8,12))
assert len(abstract.split())<=350,len(abstract.split())
assert len(binding_legend.split())<=300
# Keep only structured abstract labels bold after paragraph-level text replacement.
for idx in range(8,12):
 p=orig[idx];txt=p.text;label,body=txt.split(':',1)
 for ch in list(p._p):
  if ch.tag!=qn('w:pPr'):p._p.remove(ch)
 p.add_run(label+':').bold=True;p.add_run(body).bold=False
for p in doc.paragraphs:
 if re.match(r'^Additional file [1-6]\. ',p.text):
  txt=p.text;label,body=txt.split('. ',1)
  for ch in list(p._p):
   if ch.tag!=qn('w:pPr'):p._p.remove(ch)
  p.add_run(label+'. ').bold=True;p.add_run(body).bold=False
# Recover source scientific superscripts and subscripts after text replacements.
math_flags=[]
for el in doc.element.body.iter(qn('w:p')):
 new=text_of(el);old=original_nodes.get(el,'');flags=[None]*len(new)
 if el in original_math:
  for block in difflib.SequenceMatcher(a=old,b=new,autojunk=False).get_matching_blocks():
   flags[block.b:block.b+block.size]=original_math[el][block.a:block.a+block.size]
 restore_math(el,flags);math_flags.append(flags)
doc.core_properties.title=orig[0].text
# Track old text for each paragraph after clean document modifications.
clean_nodes=list(doc.element.body.iter(qn('w:p')))
oldtext=[original_nodes.get(el,None) for el in clean_nodes]
doc.core_properties.author='Shaoxin Huang';doc.core_properties.last_modified_by='Shaoxin Huang'
cleanpath=OUT/'Revised_Manuscript_Clean.docx';doc.save(cleanpath)
# True run-level tracked revisions, with words and punctuation preserved.
red=Document(cleanpath);rid=10000
for p_el,old in zip(red.element.body.iter(qn('w:p')),oldtext):
 new=text_of(p_el)
 if old==new:continue
 runs=p_el.findall(qn('w:r'));rpr=deepcopy(runs[0].find(qn('w:rPr'))) if runs and runs[0].find(qn('w:rPr')) is not None else None
 for x in list(p_el):
  if x.tag!=qn('w:pPr'):p_el.remove(x)
 def add(txt,kind=None):
  global rid
  if not txt:return
  r=OxmlElement('w:r')
  if rpr is not None:
   localpr=deepcopy(rpr)
   if new.startswith(('Background:','Methods:','Results:','Conclusions:','Additional file ')):
    for child in list(localpr):
     if child.tag in (qn('w:b'),qn('w:bCs')):localpr.remove(child)
   r.append(localpr)
  t=OxmlElement('w:delText' if kind=='del' else 'w:t');t.set(qn('xml:space'),'preserve');t.text=txt;r.append(t)
  if kind:
   wr=OxmlElement('w:'+kind);wr.set(qn('w:id'),str(rid));rid+=1;wr.set(qn('w:author'),'Shaoxin Huang');wr.set(qn('w:date'),'2026-10-02T08:00:00Z');wr.append(r);p_el.append(wr)
  else:p_el.append(r)
 if old is None:add(new,'ins');continue
 a=re.findall(r'\s+|\w+|[^\w\s]',old);b=re.findall(r'\s+|\w+|[^\w\s]',new)
 for tag,i,j,k,l in difflib.SequenceMatcher(a=a,b=b,autojunk=False).get_opcodes():
  if tag=='equal':add(''.join(b[k:l]))
  else:
   if i<j:add(''.join(a[i:j]),'del')
   if k<l:add(''.join(b[k:l]),'ins')
settings=red.settings.element
if settings.find(qn('w:trackRevisions')) is None:settings.append(OxmlElement('w:trackRevisions'))
for el,flags in zip(red.element.body.iter(qn('w:p')),math_flags):restore_math(el,flags)
red.save(OUT/'Revised_Manuscript_Tracked.docx')
# Verify accepting tracked changes recovers the clean text exactly.
def accepted(p):
 out=[]
 for t in p.iter(qn('w:t')):
  if not any(a.tag==qn('w:del') for a in t.iterancestors()):out.append(t.text or '')
 return ''.join(out)
assert [accepted(el) for el in red.element.body.iter(qn('w:p'))]==[text_of(el) for el in doc.element.body.iter(qn('w:p'))]
# Source Results kept verbatim except figure/citation remapping; check every original result paragraph remains.
retention=[]
for i in range(65,102):
 retention.append({'source_paragraph':i,'retained':orig[i]._p.getparent() is not None,'final_text':orig[i].text})
assert all(r['retained'] for r in retention)
(ROOT/'work/build/manuscript_change_log.json').write_text(json.dumps(changes,indent=2,ensure_ascii=False))
(ROOT/'work/build/reference_number_map.json').write_text(json.dumps(renum,indent=2))
(ROOT/'work/build/original_results_retention.json').write_text(json.dumps(retention,indent=2,ensure_ascii=False))
(ROOT/'work/build/main_build_metrics.json').write_text(json.dumps({'abstract_words':len(abstract.split()),'references':len(refs),'original_main_panels':68,'new_main_panels':6,'original_supp_panels':97,'new_supp_panels':9,'tracked_revision_elements':rid-10000,'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest()},indent=2))
print(json.dumps({'clean':str(cleanpath),'abstract_words':len(abstract.split()),'references':len(refs),'tracked_elements':rid-10000}))
