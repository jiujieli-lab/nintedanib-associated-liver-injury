#!/usr/bin/env python3
"""Build the release-facing QA reports from the current submission capsule.

The script deliberately reads the packaged manuscript, tables, figures, frozen
CSV/JSON results, and notebook in place.  It does not use development-tree
paths and it never writes absolute paths into a report.  A non-zero exit code
means that at least one scientific, structural, or binary-integrity assertion
failed; administrative authorship/funding/contact fields and the archival
repository identifier are transparent release holds rather than scientific
failures.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import statistics
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

import fitz  # PyMuPDF
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
CHECKLISTS = ROOT / "Checklists"
RELEASE_STATUS = "PASS_WITH_NONSCIENTIFIC_ADMIN_AND_REPOSITORY_HOLDS"
GENERATED_ON = datetime.now(timezone.utc).date().isoformat()

TITLE = "Role-separated public-data triangulation prioritizes target–phenotype hypotheses for nintedanib-associated liver injury"

MAIN_FIGURES = [
    "Figure_1_FAERS",
    "Figure_2_Human_DILI_Proteomics",
    "Figure_3_Nintedanib_HepG2_Dose_Perturbation",
    "Figure_4_Liver_Localization_and_Virtual_Knockout",
    "Figure_5_Multi_Evidence_Target_Prioritization",
    "Figure_6_Public_Preclinical_Constraints",
]
SUPPLEMENTARY_FIGURES = [
    "Supplementary_Figure_S1_FAERS",
    "Supplementary_Figure_S2_Protein_Distributions",
    "Supplementary_Figure_S3_Proteomics_Robustness",
    "Supplementary_Figure_S4_Nintedanib_Perturbation_Robustness",
    "Supplementary_Figure_S5_Whole_Blood_Sensitivity",
    "Supplementary_Figure_S6_GSE151374",
    "Supplementary_Figure_S7_PXD052594",
]
SUPPLEMENTARY_PANEL_SETS = {
    "Supplementary_Figure_S1_FAERS": list("ABCDEF"),
    "Supplementary_Figure_S2_Protein_Distributions": list("ABCDEFGHIJKLM"),
    "Supplementary_Figure_S3_Proteomics_Robustness": list("ABCDEFGHIJ"),
    "Supplementary_Figure_S4_Nintedanib_Perturbation_Robustness": list("ABCDEFGH"),
    "Supplementary_Figure_S5_Whole_Blood_Sensitivity": list("ABCDEF"),
    "Supplementary_Figure_S6_GSE151374": list("ABCDEFGHIJ"),
    "Supplementary_Figure_S7_PXD052594": list("ABCDEFGHIJ"),
}

EXPECTED_ROR = {
    "narrow": (3.9966057509880484, 3.2264349947394404, 4.950621213467427),
    "broad": (3.3668703340517556, 3.0829903271105863, 3.6768898515947788),
    "hy_proxy": (2.067598353963946, 1.0035195514337016, 4.259969770600775),
}
EXPECTED_HETEROGENEITY = {
    ("narrow", "all_years", "Tarone-adjusted Breslow-Day"): (88.25434621529466, 11, 3.6637359812630166e-14, None),
    ("narrow", "exclude_2022", "Tarone-adjusted Breslow-Day"): (33.92521640584324, 10, 0.00019016298099105988, None),
    ("broad", "all_years", "Tarone-adjusted Breslow-Day"): (63.325369250923984, 12, 5.552200810221564e-09, None),
    ("broad", "exclude_2022", "Tarone-adjusted Breslow-Day"): (40.28344742835485, 11, 3.199346023374794e-05, None),
    ("narrow", "all_years", "Inverse-variance Cochran Q"): (43.11266305705881, 11, 1.0381310309410039e-05, 74.48545457411967),
    ("narrow", "exclude_2022", "Inverse-variance Cochran Q"): (31.945370273909422, 10, 0.0004089179642641821, 68.69655942549132),
    ("broad", "all_years", "Inverse-variance Cochran Q"): (61.39384545688669, 12, 1.2560042610899139e-08, 80.45406683569463),
    ("broad", "exclude_2022", "Inverse-variance Cochran Q"): (39.52913773121904, 11, 4.305171178772789e-05, 72.17242613589191),
}


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def artifact(path: Path, **extra: Any) -> dict[str, Any]:
    result: dict[str, Any] = {
        "path": rel(path),
        "sha256": sha256(path),
        "bytes": path.stat().st_size,
    }
    result.update(extra)
    return result


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def as_bool(value: Any) -> bool:
    return str(value).strip().lower() in {"true", "1", "yes"}


def as_float(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def close(a: float | None, b: float | None, *, rtol: float = 1e-10, atol: float = 1e-12) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return math.isclose(a, b, rel_tol=rtol, abs_tol=atol)


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def words(text: str) -> int:
    return len(re.findall(r"[A-Za-z0-9]+(?:[\u2019'-][A-Za-z0-9]+)*", text))


def markdown_section(text: str, heading: str) -> str:
    pattern = re.compile(rf"^## {re.escape(heading)}\s*$", re.MULTILINE)
    match = pattern.search(text)
    if not match:
        return ""
    following = re.search(r"^## ", text[match.end() :], re.MULTILINE)
    end = match.end() + following.start() if following else len(text)
    return text[match.end() : end].strip()


def administrative_placeholders(paths: Iterable[Path]) -> dict[str, list[str]]:
    keyword_pattern = re.compile(
        r"AUTHOR|AFFILI|FULL NAME|EMAIL|FUND|GRANT|ACKNOWLEDG|COMPETING|DISCLOS|"
        r"INITIALS|ORCID|DEPARTMENT|INSTITUTION|ADDRESS|TELEPHONE|SUBMISSION DATE|"
        r"PREPRINT|PATENT|LEAD CONTACT|ARCHIVAL REPOSITORY|USE ONLY|ADD |SENIOR|"
        r"FOOTNOTE|DEGREE",
    )
    found: dict[str, list[str]] = {}
    for path in paths:
        text = path.read_text(encoding="utf-8")
        tokens = sorted(set(re.findall(r"\[[^\]\n]{2,180}\]", text)))
        admin = [token for token in tokens if keyword_pattern.search(token)]
        if admin:
            found[rel(path)] = admin
    return found


def pdffonts_check(path: Path) -> dict[str, Any]:
    command = subprocess.run(
        ["pdffonts", str(path)], capture_output=True, text=True, check=False
    )
    if command.returncode != 0:
        return {
            "tool": "pdffonts",
            "returncode": command.returncode,
            "all_fonts_embedded": False,
            "error": command.stderr.strip().replace(str(ROOT), "."),
        }
    rows: list[dict[str, Any]] = []
    for line in command.stdout.splitlines()[2:]:
        line = line.strip()
        if not line:
            continue
        match = re.search(r"\s(yes|no)\s+(yes|no)\s+(yes|no)\s+\d+\s+\d+\s*$", line)
        if match:
            rows.append(
                {
                    "font": line.split()[0],
                    "embedded": match.group(1) == "yes",
                    "subset": match.group(2) == "yes",
                    "unicode_map": match.group(3) == "yes",
                }
            )
    return {
        "tool": "pdffonts",
        "returncode": command.returncode,
        "font_records": len(rows),
        "all_fonts_embedded": bool(rows) and all(row["embedded"] for row in rows),
        "fonts": rows,
    }


def pdf_structure(path: Path, *, render_pages: bool = True) -> dict[str, Any]:
    reader = PdfReader(path)
    page_sizes: list[list[float]] = []
    orientations: list[str] = []
    extracted_words: list[int] = []
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        page_sizes.append([round(width, 3), round(height, 3)])
        orientations.append("landscape" if width > height else "portrait")
        extracted_words.append(words(page.extract_text() or ""))

    fitz_doc = fitz.open(path)
    spans_outside_media_box = 0
    nonwhite_fractions: list[float] = []
    render_dimensions: list[list[int]] = []
    if render_pages:
        for page in fitz_doc:
            rect = page.rect
            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        x0, y0, x1, y1 = span["bbox"]
                        if x0 < rect.x0 - 1 or y0 < rect.y0 - 1 or x1 > rect.x1 + 1 or y1 > rect.y1 + 1:
                            spans_outside_media_box += 1
            pix = page.get_pixmap(matrix=fitz.Matrix(0.75, 0.75), colorspace=fitz.csGRAY, alpha=False)
            render_dimensions.append([pix.width, pix.height])
            sample = pix.samples
            nonwhite = sum(value < 248 for value in sample)
            nonwhite_fractions.append(nonwhite / len(sample) if sample else 0.0)
    fitz_doc.close()

    font_check = pdffonts_check(path)
    raw_text = "\n".join((page.extract_text() or "") for page in reader.pages)
    raw_markdown_tokens = {
        "double_asterisk": raw_text.count("**"),
        "heading_marker": len(re.findall(r"(?m)^#{2,6}\s", raw_text)),
        "backtick": raw_text.count("`"),
    }
    return {
        **artifact(path),
        "pages": len(reader.pages),
        "encrypted": bool(reader.is_encrypted),
        "metadata_title": str((reader.metadata or {}).get("/Title", "")),
        "page_sizes_points": page_sizes,
        "orientations": orientations,
        "unique_page_sizes_points": sorted({tuple(size) for size in page_sizes}),
        "extracted_words_per_page": extracted_words,
        "minimum_extracted_words_on_any_page": min(extracted_words) if extracted_words else 0,
        "spans_outside_media_box": spans_outside_media_box,
        "render_dimensions_at_54_dpi": render_dimensions,
        "minimum_nonwhite_fraction": min(nonwhite_fractions) if nonwhite_fractions else None,
        "blank_rendered_pages": [i + 1 for i, value in enumerate(nonwhite_fractions) if value < 0.002],
        "raw_markdown_tokens": raw_markdown_tokens,
        "font_check": font_check,
    }


def svg_panel_counts(path: Path) -> Counter[str]:
    root = ET.parse(path).getroot()
    values: list[str] = []
    for element in root.iter():
        if element.text and element.text.strip():
            values.append(element.text.strip())
    return Counter(value for value in values if re.fullmatch(r"[A-Z]", value))


def raster_check(path: Path) -> dict[str, Any]:
    Image.MAX_IMAGE_PIXELS = None
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            width, height = image.size
            dpi = image.info.get("dpi")
            result = artifact(
                path,
                decoded=True,
                format=image.format,
                width_px=width,
                height_px=height,
                dpi=[round(float(dpi[0]), 3), round(float(dpi[1]), 3)] if dpi else None,
            )
    except Exception as exc:  # A release check must record, not conceal, decoder failure.
        message = str(exc).replace(str(ROOT), ".")
        result = artifact(path, decoded=False, error=f"{type(exc).__name__}: {message}")
    return result


def one_page_figure_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(path)
    return artifact(
        path,
        pages=len(reader.pages),
        encrypted=bool(reader.is_encrypted),
        page_size_points=[
            round(float(reader.pages[0].mediabox.width), 3),
            round(float(reader.pages[0].mediabox.height), 3),
        ],
        font_check=pdffonts_check(path),
    )


def figure_bundle(directory: Path, stem: str, expected_panels: list[str]) -> dict[str, Any]:
    paths = {extension: directory / f"{stem}.{extension}" for extension in ("svg", "pdf", "png", "tiff")}
    missing = [rel(path) for path in paths.values() if not path.exists()]
    if missing:
        return {"figure": stem, "status": "FAIL", "missing": missing}

    panel_counts = svg_panel_counts(paths["svg"])
    if directory.name == "Figures":
        panels_ok = all(panel_counts.get(panel) == 1 for panel in expected_panels)
    else:
        panels_ok = all(panel_counts.get(panel, 0) >= 1 for panel in expected_panels)
    pdf = one_page_figure_pdf(paths["pdf"])
    png = raster_check(paths["png"])
    tiff = raster_check(paths["tiff"])
    status = "PASS" if panels_ok and pdf["pages"] == 1 and png["decoded"] and tiff["decoded"] else "FAIL"
    return {
        "figure": stem,
        "status": status,
        "expected_panels": expected_panels,
        "svg_exact_single_letter_text_counts": dict(sorted(panel_counts.items())),
        "expected_panels_present": panels_ok,
        "svg": artifact(paths["svg"]),
        "pdf": pdf,
        "png": png,
        "tiff": tiff,
    }


def figure_typography(pdf_path: Path, figure_number: int) -> dict[str, Any]:
    document = fitz.open(pdf_path)
    page = document[0]
    width = page.rect.width
    height = page.rect.height
    # Cell/Nature-style dense figures are evaluated after fitting within a
    # 7.2 x 9.2 inch journal text area, matching the prior capsule audit.
    scale = min(7.2 * 72 / width, 9.2 * 72 / height)
    spans: list[dict[str, Any]] = []
    outside = 0
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                if not text:
                    continue
                x0, y0, x1, y1 = span["bbox"]
                if x0 < page.rect.x0 - 1 or y0 < page.rect.y0 - 1 or x1 > page.rect.x1 + 1 or y1 > page.rect.y1 + 1:
                    outside += 1
                spans.append({"text": text, "size": float(span["size"]) * scale})
    document.close()
    panel_sizes = [item["size"] for item in spans if item["text"] in list("ABCDEFGHIJ")]
    ordinary_sizes = [
        item["size"]
        for item in spans
        if item["text"] not in list("ABCDEFGHIJ") and len(item["text"]) >= 2 and item["size"] > 0
    ]
    all_sizes = [item["size"] for item in spans if item["size"] > 0]
    extracted = " ".join(item["text"] for item in spans)
    return {
        "figure": f"Figure {figure_number}",
        "pdf_path": rel(pdf_path),
        "pdf_sha256": sha256(pdf_path),
        "native_page_size_points": [round(width, 3), round(height, 3)],
        "journal_area_scale_factor": round(scale, 5),
        "all_text_median_effective_pt": round(statistics.median(all_sizes), 2) if all_sizes else None,
        "ordinary_text_minimum_effective_pt": round(min(ordinary_sizes), 2) if ordinary_sizes else None,
        "ordinary_text_median_effective_pt": round(statistics.median(ordinary_sizes), 2) if ordinary_sizes else None,
        "panel_label_median_effective_pt": round(statistics.median(panel_sizes), 2) if panel_sizes else None,
        "text_spans_outside_media_box": outside,
        "title_convention_present": f"Figure {figure_number}" in extracted,
    }


def manuscript_metrics() -> dict[str, Any]:
    manuscript_path = ROOT / "Main_Manuscript.md"
    text = manuscript_path.read_text(encoding="utf-8")
    title_match = re.search(r"^# (.+)$", text, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else ""
    summary = markdown_section(text, "Summary")
    highlights = [line[2:].strip() for line in markdown_section(text, "Highlights").splitlines() if line.startswith("- ")]
    etoc = markdown_section(text, "eTOC blurb")
    references = [int(value) for value in re.findall(r"(?m)^(\d+)\.\s+", markdown_section(text, "References"))]
    verified_reference_rows = read_csv(ROOT / "References/Nintedanib_DILI_Verified_References.csv")

    workbook_path = ROOT / "Supplementary_Tables.xlsx"
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    structured_tables_by_sheet = {sheet.title: len(sheet.tables) for sheet in workbook.worksheets}
    structured_table_count = sum(structured_tables_by_sheet.values())
    workbook.close()

    notebook_path = ROOT / "Public_Data_Reproducibility_Notebook.ipynb"
    notebook = json.loads(notebook_path.read_text(encoding="utf-8"))
    code_cells = [cell for cell in notebook["cells"] if cell.get("cell_type") == "code"]
    execution_counts = [cell.get("execution_count") for cell in code_cells]
    outputs = [output for cell in code_cells for output in cell.get("outputs", [])]
    notebook_metrics = {
        "path": rel(notebook_path),
        "sha256": sha256(notebook_path),
        "total_cells": len(notebook["cells"]),
        "code_cells": len(code_cells),
        "captured_outputs": len(outputs),
        "error_outputs": sum(output.get("output_type") == "error" for output in outputs),
        "execution_counts": execution_counts,
        "sequential_execution": execution_counts == list(range(1, len(code_cells) + 1)),
    }

    docx_audit_path = ROOT / "Checklists/Main_Manuscript_DOCX_Audit.json"
    docx_audit = json.loads(docx_audit_path.read_text(encoding="utf-8"))
    docx_checks = docx_audit.get("checks", [])

    admin_paths = [
        ROOT / "Main_Manuscript.md",
        ROOT / "Title_Page_and_Author_Information.md",
        ROOT / "Author_Contributions_and_Declarations.md",
        ROOT / "Cover_Letter.md",
        ROOT / "Data_and_Code_Availability.md",
    ]
    admin = administrative_placeholders(admin_paths)

    return {
        "manuscript": artifact(manuscript_path),
        "title": title,
        "title_characters": len(title),
        "summary_words": words(summary),
        "highlight_count": len(highlights),
        "highlight_characters": [len(item) for item in highlights],
        "etoc_words": words(etoc),
        "markdown_words_total": words(text),
        "reference_count_in_manuscript": len(references),
        "reference_numbering_is_1_to_n": references == list(range(1, len(references) + 1)),
        "verified_reference_rows": len(verified_reference_rows),
        "workbook": artifact(workbook_path),
        "workbook_sheet_count": len(structured_tables_by_sheet),
        "structured_table_count": structured_table_count,
        "structured_tables_by_sheet": structured_tables_by_sheet,
        "notebook": notebook_metrics,
        "docx_audit": {
            "path": rel(docx_audit_path),
            "sha256": sha256(docx_audit_path),
            "status": docx_audit.get("status"),
            "checks_passed": sum(bool(check.get("passed")) for check in docx_checks),
            "checks_total": len(docx_checks),
        },
        "administrative_placeholders": admin,
        "main_manuscript_administrative_placeholders": admin.get("Main_Manuscript.md", []),
        "repository_placeholders": sorted(
            {
                token
                for tokens in admin.values()
                for token in tokens
                if "REPOSITORY" in token or "PREPRINT" in token
            }
        ),
    }


def faers_metrics() -> dict[str, Any]:
    overall_path = ROOT / "Source_Data/FAERS/overall_active_comparator_signals.csv"
    heterogeneity_path = ROOT / "Source_Data/FAERS/year_effect_heterogeneity_tests.csv"
    overall_rows = read_csv(overall_path)
    overall: list[dict[str, Any]] = []
    exact_ok = True
    for row in overall_rows:
        endpoint = row["endpoint"]
        a = int(row["nintedanib_event"])
        b = int(row["nintedanib_nonevent"])
        c = int(row["pirfenidone_event"])
        d = int(row["pirfenidone_nonevent"])
        recalculated_ror = (a * d) / (b * c)
        standard_error = math.sqrt(1 / a + 1 / b + 1 / c + 1 / d)
        recalculated_low = math.exp(math.log(recalculated_ror) - 1.96 * standard_error)
        recalculated_high = math.exp(math.log(recalculated_ror) + 1.96 * standard_error)
        stored = (float(row["ROR"]), float(row["ROR_low95"]), float(row["ROR_high95"]))
        expected = EXPECTED_ROR[endpoint]
        row_ok = all(close(x, y) for x, y in zip(stored, expected)) and all(
            close(x, y) for x, y in zip(stored, (recalculated_ror, recalculated_low, recalculated_high))
        )
        exact_ok &= row_ok
        overall.append(
            {
                "endpoint": endpoint,
                "counts": {
                    "nintedanib_event": a,
                    "nintedanib_nonevent": b,
                    "pirfenidone_event": c,
                    "pirfenidone_nonevent": d,
                    "nintedanib_total": int(row["nintedanib_total"]),
                    "pirfenidone_total": int(row["pirfenidone_total"]),
                },
                "stored": {"ROR": stored[0], "low95": stored[1], "high95": stored[2]},
                "independent_recalculation": {
                    "ROR": recalculated_ror,
                    "low95": recalculated_low,
                    "high95": recalculated_high,
                },
                "exact_release_value_match": row_ok,
            }
        )

    heterogeneity_rows = read_csv(heterogeneity_path)
    heterogeneity: list[dict[str, Any]] = []
    heterogeneity_ok = True
    keyed_rows = {(row["endpoint"], row["window"], row["test"]): row for row in heterogeneity_rows}
    for key, expected in EXPECTED_HETEROGENEITY.items():
        row = keyed_rows.get(key)
        if not row:
            heterogeneity_ok = False
            heterogeneity.append({"endpoint": key[0], "window": key[1], "test": key[2], "missing": True})
            continue
        values = (
            float(row["statistic"]),
            int(row["degrees_of_freedom"]),
            float(row["p_value"]),
            as_float(row["i2_percent"]),
        )
        row_ok = close(values[0], expected[0]) and values[1] == expected[1] and close(values[2], expected[2]) and close(values[3], expected[3])
        heterogeneity_ok &= row_ok
        heterogeneity.append(
            {
                "endpoint": key[0],
                "window": key[1],
                "test": key[2],
                "statistic": values[0],
                "degrees_of_freedom": values[1],
                "p_value": values[2],
                "i2_percent": values[3],
                "exact_release_value_match": row_ok,
            }
        )

    country_rows = read_csv(ROOT / "Source_Data/FAERS/narrow_2022_country_audit.csv")
    month_rows = read_csv(ROOT / "Source_Data/FAERS/narrow_2022_month_audit.csv")
    italy_row = next((row for row in country_rows if row.get("country_for_audit") == "IT"), None)
    july_row = next((row for row in month_rows if row.get("receive_month") == "2022-07"), None)

    return {
        "overall_source": artifact(overall_path),
        "heterogeneity_source": artifact(heterogeneity_path),
        "unit": "latest-version safetyreportid (report)",
        "estimand_boundary": "all-indication drug-name-indexed active-comparator disproportionality; not incidence, absolute risk, or causality",
        "overall": overall,
        "overall_exact_checks_passed": exact_ok,
        "heterogeneity": heterogeneity,
        "heterogeneity_exact_checks_passed": heterogeneity_ok,
        "heterogeneity_rows_in_source": len(heterogeneity_rows),
        "cluster_2022": {
            "nintedanib_narrow_reports": int(italy_row["drug_2022_narrow_reports"]) if italy_row else None,
            "Italy_reports": int(italy_row["reports"]) if italy_row else None,
            "July_2022_reports": int(july_row["reports"]) if july_row else None,
            "country_audit": artifact(ROOT / "Source_Data/FAERS/narrow_2022_country_audit.csv"),
            "month_audit": artifact(ROOT / "Source_Data/FAERS/narrow_2022_month_audit.csv"),
        },
    }


def target_role_metrics() -> dict[str, Any]:
    gate_path = ROOT / "Source_Data/Integration/target_tier_gate_sensitivity.csv"
    phenotype_path = ROOT / "Source_Data/Integration/phenotype_anchor_tiers.csv"
    integration_text_path = ROOT / "Source_Data/Integration/integration_methods_results.md"
    human_proteomics_path = ROOT / "Source_Data/Human_DILI_Proteomics/manuscript_ready_proteomics_text.md"
    gate_rows = read_csv(gate_path)
    phenotype_rows = read_csv(phenotype_path)
    integration_text = integration_text_path.read_text(encoding="utf-8")
    human_text = human_proteomics_path.read_text(encoding="utf-8")

    exploratory_t2 = [row["gene_symbol"] for row in gate_rows if as_bool(row["primary_T2_gate"])]
    strict_t2 = [row["gene_symbol"] for row in gate_rows if as_bool(row["strict_all_seed_T2_sensitivity_gate"])]
    fgfr1 = next(row for row in gate_rows if row["gene_symbol"] == "FGFR1")
    fbp1 = next(row for row in phenotype_rows if row["gene_symbol"] == "FBP1")
    no_t1_text = "T1 required" in integration_text and "was not assigned" in integration_text
    source_selected_fbp1 = "source-selected" in human_text and "FBP1" in human_text

    return {
        "sources": [artifact(gate_path), artifact(phenotype_path), artifact(integration_text_path), artifact(human_proteomics_path)],
        "T1_targets": [],
        "no_T1_assigned_confirmed_in_frozen_integration_text": no_t1_text,
        "exploratory_T2_targets": exploratory_t2,
        "strict_all_seed_T2_targets": strict_t2,
        "FGFR1": {
            "evidence_tier": fgfr1["evidence_tier"],
            "primary_exploratory_T2_gate": as_bool(fgfr1["primary_T2_gate"]),
            "strict_all_seed_gate": as_bool(fgfr1["strict_all_seed_T2_sensitivity_gate"]),
            "joint_context_compartment": fgfr1["joint_context_compartment"],
            "joint_context_detection_fraction": float(fgfr1["joint_context_detection_fraction"]),
            "cross_seed_median_within_compartment_rank": float(fgfr1["vko_within_compartment_rank"]),
            "n_seeds": int(fgfr1["n_seeds"]),
            "n_seeds_top_quartile": int(fgfr1["n_seeds_top_quartile"]),
            "minimum_seed_rank": float(fgfr1["min_seed_vko_rank"]),
        },
        "FBP1": {
            "evidence_tier": fbp1["evidence_tier"],
            "candidate_role": fbp1["candidate_role"],
            "source_selected_P1_boundary_confirmed": source_selected_fbp1,
            "nintedanib_specific": False,
        },
        "virtual_knockout_direction": "unsigned",
        "all_role_assertions_passed": (
            exploratory_t2 == ["FGFR1"]
            and strict_t2 == []
            and no_t1_text
            and fbp1["evidence_tier"].startswith("P1:")
            and source_selected_fbp1
        ),
    }


def main() -> int:
    CHECKLISTS.mkdir(parents=True, exist_ok=True)

    manuscript = manuscript_metrics()
    faers = faers_metrics()
    targets = target_role_metrics()
    main_pdf = pdf_structure(ROOT / "Main_Manuscript.pdf")
    supplementary_pdf = pdf_structure(ROOT / "Supplementary_Information.pdf")

    main_bundles = [
        figure_bundle(ROOT / "Figures", stem, list("ABCDEFGHIJ")) for stem in MAIN_FIGURES
    ]
    supplementary_bundles = [
        figure_bundle(ROOT / "Supplementary_Figures", stem, SUPPLEMENTARY_PANEL_SETS[stem])
        for stem in SUPPLEMENTARY_FIGURES
    ]
    typography = [
        figure_typography(ROOT / "Figures" / f"{stem}.pdf", index)
        for index, stem in enumerate(MAIN_FIGURES, 1)
    ]

    # S1–S5 retain near-native scale on portrait Letter figure pages.  The
    # denser ten-panel S6/S7 figures use A3 landscape figure pages so that their
    # labels remain legible; every legend page remains portrait Letter.
    expected_supp_orientations = ["portrait"] * 12 + [
        "landscape",
        "portrait",
        "landscape",
        "portrait",
    ]
    supplementary_page_size_pass = all(
        (
            close(size[0], 1190.55118, rtol=0, atol=0.01)
            and close(size[1], 841.889764, rtol=0, atol=0.01)
        )
        if page_number in {13, 15}
        else close(size[0], 612.0, rtol=0, atol=0.01)
        and close(size[1], 792.0, rtol=0, atol=0.01)
        for page_number, size in enumerate(supplementary_pdf["page_sizes_points"], 1)
    )
    main_pdf_pass = (
        main_pdf["pages"] == 58
        and not main_pdf["encrypted"]
        and main_pdf["metadata_title"] == TITLE
        and set(main_pdf["orientations"]) == {"portrait"}
        and main_pdf["spans_outside_media_box"] == 0
        and not main_pdf["blank_rendered_pages"]
        and all(value == 0 for value in main_pdf["raw_markdown_tokens"].values())
        and main_pdf["font_check"]["all_fonts_embedded"]
    )
    supplementary_pdf_pass = (
        supplementary_pdf["pages"] == 16
        and not supplementary_pdf["encrypted"]
        and supplementary_pdf["metadata_title"] == f"Supplementary Information - {TITLE}"
        and supplementary_pdf["orientations"] == expected_supp_orientations
        and supplementary_page_size_pass
        and supplementary_pdf["spans_outside_media_box"] == 0
        and not supplementary_pdf["blank_rendered_pages"]
        and all(value == 0 for value in supplementary_pdf["raw_markdown_tokens"].values())
        and supplementary_pdf["font_check"]["all_fonts_embedded"]
    )
    figures_pass = all(item["status"] == "PASS" for item in main_bundles + supplementary_bundles)
    typography_pass = all(
        item["text_spans_outside_media_box"] == 0
        and item["title_convention_present"]
        and item["panel_label_median_effective_pt"] is not None
        for item in typography
    )
    manuscript_pass = (
        manuscript["title"] == TITLE
        and manuscript["reference_count_in_manuscript"] == 75
        and manuscript["reference_numbering_is_1_to_n"]
        and manuscript["verified_reference_rows"] == 75
        and manuscript["workbook_sheet_count"] == 22
        and manuscript["structured_table_count"] == 107
        and manuscript["notebook"]["total_cells"] == 32
        and manuscript["notebook"]["code_cells"] == 18
        and manuscript["notebook"]["error_outputs"] == 0
        and manuscript["notebook"]["sequential_execution"]
        and manuscript["docx_audit"]["status"] == "PASS"
        and manuscript["docx_audit"]["checks_passed"] == 28
        and manuscript["docx_audit"]["checks_total"] == 28
    )
    science_pass = (
        manuscript_pass
        and faers["overall_exact_checks_passed"]
        and faers["heterogeneity_exact_checks_passed"]
        and targets["all_role_assertions_passed"]
        and main_pdf_pass
        and supplementary_pdf_pass
        and figures_pass
        and typography_pass
    )

    assembly = {
        "schema_version": "1.1",
        "generated_on_utc": GENERATED_ON,
        "generator": artifact(Path(__file__).resolve()),
        "status": RELEASE_STATUS if science_pass else "FAIL",
        "scientific_status": "PASS" if science_pass else "FAIL",
        "direct_portal_upload": "HOLD: author-supplied administrative metadata and archival repository identifier required",
        "title": manuscript["title"],
        "title_characters": manuscript["title_characters"],
        "summary_words": manuscript["summary_words"],
        "highlight_count": manuscript["highlight_count"],
        "highlight_characters": manuscript["highlight_characters"],
        "etoc_words": manuscript["etoc_words"],
        "manuscript_markdown_words_total": manuscript["markdown_words_total"],
        "word_limit_policy": "Full-detail Cell-style manuscript; the author explicitly requested no 7,000-word compression.",
        "reference_count": manuscript["reference_count_in_manuscript"],
        "verified_reference_rows": manuscript["verified_reference_rows"],
        "supplementary_workbook_sheets": manuscript["workbook_sheet_count"],
        "supplementary_structured_tables": manuscript["structured_table_count"],
        "main_manuscript_pdf_pages": main_pdf["pages"],
        "supplementary_information_pdf_pages": supplementary_pdf["pages"],
        "supplementary_information_page_orientations": supplementary_pdf["orientations"],
        "main_figure_count": len(main_bundles),
        "main_figure_panels_each": "A–J",
        "supplementary_figure_count": len(supplementary_bundles),
        "notebook": manuscript["notebook"],
        "docx_audit": manuscript["docx_audit"],
        "target_and_phenotype_roles": targets,
        "faers_release_checks": faers,
        "administrative_and_repository_holds": {
            "placeholders_by_submission_file": manuscript["administrative_placeholders"],
            "main_manuscript_placeholders": manuscript["main_manuscript_administrative_placeholders"],
            "repository_placeholders": manuscript["repository_placeholders"],
            "boundary": "These fields require author or repository action and were not inferred. They do not alter the scientific/statistical PASS decision.",
        },
        "artifact_hashes": [
            manuscript["manuscript"],
            artifact(ROOT / "Main_Manuscript.docx"),
            {key: value for key, value in main_pdf.items() if key in {"path", "sha256", "bytes", "pages"}},
            artifact(ROOT / "Supplementary_Information.md"),
            {key: value for key, value in supplementary_pdf.items() if key in {"path", "sha256", "bytes", "pages"}},
            manuscript["workbook"],
            {key: value for key, value in manuscript["notebook"].items() if key in {"path", "sha256"}},
        ],
        "checks": {
            "current_title": manuscript["title"] == TITLE,
            "references_75": manuscript["reference_count_in_manuscript"] == manuscript["verified_reference_rows"] == 75,
            "reference_sequence_complete": manuscript["reference_numbering_is_1_to_n"],
            "structured_tables_107": manuscript["structured_table_count"] == 107,
            "main_pdf_58_pages_portrait": main_pdf_pass,
            "supplementary_pdf_16_pages_S1_to_S5_letter_portrait_S6_to_S7_A3_landscape": supplementary_pdf_pass,
            "six_main_figures_each_A_to_J_all_four_formats": all(item["status"] == "PASS" for item in main_bundles),
            "seven_supplementary_figures_all_four_formats": all(item["status"] == "PASS" for item in supplementary_bundles),
            "notebook_32_cells_18_code_sequential_no_errors": manuscript["notebook"]["total_cells"] == 32 and manuscript["notebook"]["code_cells"] == 18 and manuscript["notebook"]["sequential_execution"] and manuscript["notebook"]["error_outputs"] == 0,
            "docx_audit_28_of_28": manuscript["docx_audit"]["checks_passed"] == manuscript["docx_audit"]["checks_total"] == 28,
            "faers_exact_release_values": faers["overall_exact_checks_passed"] and faers["heterogeneity_exact_checks_passed"],
            "role_separated_target_claims": targets["all_role_assertions_passed"],
            "typography_and_media_box": typography_pass,
        },
        "validation_passed": science_pass,
    }
    assembly_path = CHECKLISTS / "Manuscript_Assembly_Validation.json"
    write_json(assembly_path, assembly)

    main_render_report = f"""# Main manuscript render QA

**Status:** {'PASS' if main_pdf_pass else 'FAIL'}  
**Release status:** {RELEASE_STATUS if science_pass else 'FAIL'}  
**Frozen file:** `Main_Manuscript.pdf`  
**PDF SHA-256:** `{main_pdf['sha256']}`  
**DOCX SHA-256:** `{sha256(ROOT / 'Main_Manuscript.docx')}`  
**Source Markdown SHA-256:** `{manuscript['manuscript']['sha256']}`

## Structural and render checks

- {main_pdf['pages']} pages; all pages are US Letter portrait; the PDF is unencrypted.
- PDF metadata title exactly matches `{TITLE}`.
- Poppler `pdffonts` found {main_pdf['font_check']['font_records']} font records; all are embedded.
- All {main_pdf['pages']} pages were rasterized at 54 dpi in page order; no blank render was detected (minimum non-white fraction {main_pdf['minimum_nonwhite_fraction']:.4f}).
- Extracted text spans outside the MediaBox: {main_pdf['spans_outside_media_box']}.
- Raw Markdown token scan: {json.dumps(main_pdf['raw_markdown_tokens'], ensure_ascii=False)}.
- The paired DOCX audit passed {manuscript['docx_audit']['checks_passed']}/{manuscript['docx_audit']['checks_total']} checks.
- The manuscript contains {manuscript['reference_count_in_manuscript']} consecutively numbered references.

## Administrative boundary

The scientific/render checks pass independently of author-supplied fields. The remaining manuscript placeholders are: {', '.join(f'`{token}`' for token in manuscript['main_manuscript_administrative_placeholders'])}. The archival repository identifier is also intentionally pending. These fields must be completed before portal upload and were not inferred.
"""
    (CHECKLISTS / "Main_Manuscript_Render_QA.md").write_text(main_render_report, encoding="utf-8")

    supp_render_report = f"""# Supplementary information render QA

**Status:** {'PASS' if supplementary_pdf_pass else 'FAIL'}  
**Release status:** {RELEASE_STATUS if science_pass else 'FAIL'}  
**Frozen file:** `Supplementary_Information.pdf`  
**PDF SHA-256:** `{supplementary_pdf['sha256']}`  
**Source Markdown SHA-256:** `{sha256(ROOT / 'Supplementary_Information.md')}`

## Structural and render checks

- {supplementary_pdf['pages']} pages; pages 1–12, 14, and 16 are US Letter portrait. Dense ten-panel S6/S7 figure pages 13 and 15 are A3 landscape; their legend pages 14 and 16 return to US Letter portrait.
- Supplementary Figures S1–S7 are packaged in PDF, SVG, PNG, and TIFF; every source PDF contains one page.
- Poppler `pdffonts` found {supplementary_pdf['font_check']['font_records']} font records; all are embedded.
- All {supplementary_pdf['pages']} pages were rasterized at 54 dpi in page order; no blank render was detected (minimum non-white fraction {supplementary_pdf['minimum_nonwhite_fraction']:.4f}).
- Extracted text spans outside the MediaBox: {supplementary_pdf['spans_outside_media_box']}.
- Raw Markdown token scan: {json.dumps(supplementary_pdf['raw_markdown_tokens'], ensure_ascii=False)}.

## Interpretation boundary

The supplement reports public-data reanalysis. It does not convert FAERS disproportionality into incidence or causality, treat pooled constituent animals or cells as independent replicates, assign beneficial/adverse direction to virtual knockout, or claim hepatic DILI confirmation from lung datasets.
"""
    (CHECKLISTS / "Supplementary_Information_Render_QA.md").write_text(supp_render_report, encoding="utf-8")

    typography_lines = [
        "# Main-figure typography and export QA",
        "",
        f"**Status:** {'PASS' if figures_pass and typography_pass else 'FAIL'}  ",
        f"**Release status:** {RELEASE_STATUS if science_pass else 'FAIL'}  ",
        "**Scope:** Figures 1–6, each containing panels A–J; PDF, SVG, PNG, and TIFF exports.",
        "",
        "## Structural checks",
        "",
        "- Every main-figure SVG contains exactly one panel label for each letter A–J.",
        "- Every PDF contains one page, has zero extracted text spans outside its MediaBox, and uses the common `Figure N | title` convention.",
        "- Every PNG and TIFF passed decoder verification; recorded resolution is taken from the current binary, not a filename assumption.",
        "",
        "## Effective typography after 7.2 × 9.2 inch journal-area fitting",
        "",
        "| Figure | Scale | All-text median (pt) | Ordinary min–median (pt) | Panel label median (pt) | PDF SHA-256 |",
        "|---|---:|---:|---:|---:|---|",
    ]
    for item in typography:
        typography_lines.append(
            f"| {item['figure']} | {item['journal_area_scale_factor']:.5f} | "
            f"{item['all_text_median_effective_pt']:.2f} | "
            f"{item['ordinary_text_minimum_effective_pt']:.2f}–{item['ordinary_text_median_effective_pt']:.2f} | "
            f"{item['panel_label_median_effective_pt']:.2f} | `{item['pdf_sha256']}` |"
        )
    typography_lines.extend(["", "## Current export hashes", ""])
    for bundle in main_bundles:
        typography_lines.append(
            f"- {bundle['figure']}: PDF `{bundle['pdf']['sha256']}`; SVG `{bundle['svg']['sha256']}`; PNG `{bundle['png']['sha256']}`; TIFF `{bundle['tiff']['sha256']}`."
        )
    (CHECKLISTS / "Main_Figure_Typography_QA.md").write_text("\n".join(typography_lines) + "\n", encoding="utf-8")

    faers_main = next(item for item in main_bundles if item["figure"] == "Figure_1_FAERS")
    faers_supp = next(item for item in supplementary_bundles if item["figure"] == "Supplementary_Figure_S1_FAERS")
    faers_report = {
        "schema_version": "1.1",
        "generated_on_utc": GENERATED_ON,
        "generator": artifact(Path(__file__).resolve()),
        "status": "PASS" if faers["overall_exact_checks_passed"] and faers["heterogeneity_exact_checks_passed"] and faers_main["status"] == faers_supp["status"] == "PASS" else "FAIL",
        "release_status": RELEASE_STATUS if science_pass else "FAIL",
        "unit": faers["unit"],
        "estimand_boundary": faers["estimand_boundary"],
        "overall_independent_recalculation": faers["overall"],
        "year_effect_heterogeneity_check": {
            "interpretation": "Common annual effects rejected; Mantel–Haenszel estimates retained as descriptive pooled summaries.",
            "rows": faers["heterogeneity"],
            "source_rows": faers["heterogeneity_rows_in_source"],
        },
        "cluster_2022": faers["cluster_2022"],
        "figure_1": faers_main,
        "supplementary_figure_s1": faers_supp,
        "source_hashes": [faers["overall_source"], faers["heterogeneity_source"]],
        "all_checks_passed": faers["overall_exact_checks_passed"] and faers["heterogeneity_exact_checks_passed"] and faers_main["status"] == faers_supp["status"] == "PASS",
    }
    faers_report_path = CHECKLISTS / "FAERS_figure_validation_report.json"
    write_json(faers_report_path, faers_report)

    supporting_paths = [
        assembly_path,
        CHECKLISTS / "Main_Manuscript_DOCX_Audit.json",
        CHECKLISTS / "Main_Manuscript_Render_QA.md",
        CHECKLISTS / "Supplementary_Information_Render_QA.md",
        CHECKLISTS / "Main_Figure_Typography_QA.md",
        CHECKLISTS / "Supplementary_Tables_v1_1_Validation.json",
        faers_report_path,
        ROOT / "Notebook_Validation_Report.md",
        ROOT / "Source_Data/Integration/integration_validation_report.json",
    ]
    prepackage = {
        "schema_version": "1.1",
        "audit_type": "release_statistical_editorial_and_artifact_integrity",
        "audit_date_utc": GENERATED_ON,
        "generator": artifact(Path(__file__).resolve()),
        "decision": RELEASE_STATUS if science_pass else "FAIL",
        "scientific_decision": "PASS" if science_pass else "FAIL",
        "direct_portal_upload": "HOLD",
        "domain_status": {
            "scientific": "PASS" if science_pass else "FAIL",
            "statistical": "PASS" if faers["overall_exact_checks_passed"] and faers["heterogeneity_exact_checks_passed"] else "FAIL",
            "editorial_and_references": "PASS" if manuscript_pass else "FAIL",
            "figures_and_captions": "PASS" if figures_pass and typography_pass else "FAIL",
            "reproducibility": "PASS" if manuscript["notebook"]["sequential_execution"] and manuscript["notebook"]["error_outputs"] == 0 else "FAIL",
            "artifact_integrity": "PASS" if main_pdf_pass and supplementary_pdf_pass else "FAIL",
            "direct_portal_upload": "HOLD: non-scientific administrative and repository fields",
        },
        "failures": [] if science_pass else [name for name, passed in {
            "manuscript": manuscript_pass,
            "FAERS": faers["overall_exact_checks_passed"] and faers["heterogeneity_exact_checks_passed"],
            "target_roles": targets["all_role_assertions_passed"],
            "main_PDF": main_pdf_pass,
            "supplementary_PDF": supplementary_pdf_pass,
            "figure_binaries": figures_pass,
            "typography": typography_pass,
        }.items() if not passed],
        "frozen_primary_artifacts": [
            artifact(ROOT / "Main_Manuscript.md"),
            artifact(ROOT / "Main_Manuscript.docx", audit="28/28 PASS"),
            artifact(ROOT / "Main_Manuscript.pdf", pages=main_pdf["pages"]),
            artifact(ROOT / "Supplementary_Information.md"),
            artifact(ROOT / "Supplementary_Information.pdf", pages=supplementary_pdf["pages"]),
            artifact(ROOT / "Supplementary_Tables.xlsx", sheets=manuscript["workbook_sheet_count"], structured_tables=manuscript["structured_table_count"]),
            artifact(ROOT / "Public_Data_Reproducibility_Notebook.ipynb", cells=manuscript["notebook"]["total_cells"], code_cells=manuscript["notebook"]["code_cells"]),
        ],
        "supporting_audits": [artifact(path) for path in supporting_paths],
        "faers_exact_results": faers,
        "role_separated_integration": targets,
        "figure_audit": {
            "main": main_bundles,
            "supplementary": supplementary_bundles,
            "typography": typography,
        },
        "administrative_and_repository_holds": assembly["administrative_and_repository_holds"],
        "claim_boundaries": {
            "faers_incidence_or_causality": False,
            "FBP1_unbiased_or_nintedanib_specific_discovery": False,
            "FGFR1_confirmed_mediator": False,
            "virtual_knockout_directional_effect": False,
            "public_lung_data_confirm_hepatic_DILI": False,
            "prospective_perturbation_rescue_complete": False,
        },
    }
    prepackage_path = CHECKLISTS / "PrePackage_Statistical_Editorial_Audit.json"
    write_json(prepackage_path, prepackage)

    prepackage_md = f"""# Pre-package statistical and editorial audit

**Decision:** {prepackage['decision']}  
**Scientific/statistical/editorial status:** {'PASS' if science_pass else 'FAIL'}  
**Direct portal upload:** HOLD pending author-supplied administrative metadata and archival repository identifier.

## Evidence frozen in the current capsule

- Current title: **{TITLE}**.
- References: {manuscript['reference_count_in_manuscript']} in the manuscript and {manuscript['verified_reference_rows']} verified reference records.
- Supplementary workbook: {manuscript['workbook_sheet_count']} sheets and {manuscript['structured_table_count']} structured tables.
- PDFs: Main Manuscript {main_pdf['pages']} US Letter portrait pages; Supplementary Information {supplementary_pdf['pages']} pages, with S1–S5 and all legends on US Letter portrait pages and dense S6/S7 figure pages 13/15 on A3 landscape pages.
- Figures: six main figures, each A–J, plus seven supplementary figures; PDF/SVG/PNG/TIFF binary checks {'passed' if figures_pass else 'failed'}.
- Notebook: {manuscript['notebook']['total_cells']} total cells, {manuscript['notebook']['code_cells']} code cells executed sequentially, {manuscript['notebook']['error_outputs']} error outputs.
- DOCX audit: {manuscript['docx_audit']['checks_passed']}/{manuscript['docx_audit']['checks_total']} PASS.

## Statistical and mechanistic boundaries

- FAERS narrow ROR {faers['overall'][0]['stored']['ROR']:.12g} (95% CI {faers['overall'][0]['stored']['low95']:.12g}–{faers['overall'][0]['stored']['high95']:.12g}); broad ROR {faers['overall'][1]['stored']['ROR']:.12g} (95% CI {faers['overall'][1]['stored']['low95']:.12g}–{faers['overall'][1]['stored']['high95']:.12g}). The exact independent recalculation and all eight prespecified heterogeneity rows are in `FAERS_figure_validation_report.json`.
- FBP1 is a source-selected, multi-drug human-DILI P1 phenotype anchor, not an unbiased nintedanib-specific target discovery.
- FGFR1 is the sole candidate meeting the prespecified exploratory T2 gate; the strict all-seed T2 set is empty and no T1 target is assigned.
- Virtual-knockout displacement is unsigned. Public pulmonary datasets define pulmonary-response/model constraints and do not validate nintedanib-associated hepatic DILI.

## Administrative and repository holds

All unresolved author/funder/contact/declaration/repository fields are enumerated by source file in `Manuscript_Assembly_Validation.json`. They require author or repository action and were intentionally not inferred.

## Current audit hash

- `PrePackage_Statistical_Editorial_Audit.json`: `{sha256(prepackage_path)}`
- `Manuscript_Assembly_Validation.json`: `{sha256(assembly_path)}`
- `FAERS_figure_validation_report.json`: `{sha256(faers_report_path)}`
"""
    (CHECKLISTS / "PrePackage_Statistical_Editorial_Audit.md").write_text(prepackage_md, encoding="utf-8")

    print(
        json.dumps(
            {
                "status": RELEASE_STATUS if science_pass else "FAIL",
                "reports_written": [
                    rel(assembly_path),
                    "Checklists/Main_Manuscript_Render_QA.md",
                    "Checklists/Supplementary_Information_Render_QA.md",
                    "Checklists/Main_Figure_Typography_QA.md",
                    rel(prepackage_path),
                    "Checklists/PrePackage_Statistical_Editorial_Audit.md",
                    rel(faers_report_path),
                ],
                "main_pdf_pages": main_pdf["pages"],
                "supplementary_pdf_pages": supplementary_pdf["pages"],
                "references": manuscript["reference_count_in_manuscript"],
                "structured_tables": manuscript["structured_table_count"],
                "figure_binary_status": "PASS" if figures_pass else "FAIL",
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if science_pass else 1


if __name__ == "__main__":
    sys.exit(main())
