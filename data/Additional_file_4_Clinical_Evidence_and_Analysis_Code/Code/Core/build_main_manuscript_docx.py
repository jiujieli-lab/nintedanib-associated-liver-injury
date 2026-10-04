#!/usr/bin/env python3
"""Build one submission-ready DOCX from the assembled manuscript Markdown.

Design preset: narrative_proposal.
First-page header pattern: proposal_centerpiece, adapted as an academic title page.
Named overrides: AcademicTitle, ScientificTable, ReferenceEntry, FigureLegend.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "Main_Manuscript.md"
OUTPUT = ROOT / "Main_Manuscript.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
GRAY = "5B6573"
LIGHT = "F4F6F9"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360
TABLE_INDENT = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, val in CELL_MARGINS.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


CITATION = r"\[(?:\d+(?:[,–—-]\d+)*)\]"
INLINE = re.compile(
    rf"({CITATION}|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if re.fullmatch(CITATION, token):
            run = paragraph.add_run(token[1:-1])
            run.font.superscript = True
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_run_fonts(doc: Document) -> None:
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")


def configure_styles(doc: Document) -> None:
    set_run_fonts(doc)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
        "Heading 4": (11, INK, 6, 3),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for name in ["Bullet Body", "Number Body", "Reference Entry", "Figure Legend", "Scientific Table"]:
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    for name in ["Bullet Body", "Number Body"]:
        style = doc.styles[name]
        style.base_style = normal
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208
    ref = doc.styles["Reference Entry"]
    ref.base_style = normal
    ref.font.size = Pt(9.5)
    ref.paragraph_format.left_indent = Inches(0.25)
    ref.paragraph_format.first_line_indent = Inches(-0.25)
    ref.paragraph_format.space_after = Pt(4)
    ref.paragraph_format.line_spacing = 1.0
    legend = doc.styles["Figure Legend"]
    legend.base_style = normal
    # Keep the long result-figure and supplemental legends together without
    # creating a two-line terminal orphan page.  The modest 9.25-pt setting is
    # still comfortably within journal submission readability norms.
    legend.font.size = Pt(9.25)
    legend.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    legend.paragraph_format.space_after = Pt(6)
    legend.paragraph_format.line_spacing = 1.08
    table_style = doc.styles["Scientific Table"]
    table_style.base_style = normal
    table_style.font.size = Pt(8.5)
    table_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_style.paragraph_format.space_after = Pt(2)
    table_style.paragraph_format.line_spacing = 1.0


def add_numbering(doc: Document, fmt: str) -> int:
    root = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))]
    abstract_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if fmt == "bullet" else "%1.")
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "279")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, jc, ppr])
    abstract.append(lvl)
    first_num = root.find(qn("w:num"))
    if first_num is None:
        root.append(abstract)
    else:
        root.insert(list(root).index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    root.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, nid])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:distance"), "360")
    line_numbers.set(qn("w:restart"), "continuous")
    section._sectPr.append(line_numbers)
    header = section.header.paragraphs[0]
    header.text = "CELL REPORTS MEDICINE | RESEARCH ARTICLE"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_title_page(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(44)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("RESEARCH ARTICLE")
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run("Proteomics • pharmacovigilance • single-cell networks • public preclinical constraints")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    for text, bold in [
        ("[AUTHOR LIST]", True),
        ("[AFFILIATIONS]", False),
        ("Lead contact and correspondence: [FULL NAME] ([EMAIL])", False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.size = Pt(11 if bold else 10)
        run.bold = bold
        run.font.color.rgb = RGBColor.from_string(INK if bold else GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    run = p.add_run("Target journal: Cell Reports Medicine | Public-data reanalysis")
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_page_break()


def table_widths(ncols: int, rows: list[list[str]]) -> list[int]:
    if ncols == 2:
        return [2700, 6660]
    if ncols == 3:
        return [1700, 3000, 4660]
    if ncols == 4:
        return [1700, 1900, 2800, 2960]
    lengths = [max(len(row[i]) if i < len(row) else 0 for row in rows) for i in range(ncols)]
    weights = [max(1, min(x, 60)) for x in lengths]
    raw = [int(TABLE_WIDTH * x / sum(weights)) for x in weights]
    raw[-1] += TABLE_WIDTH - sum(raw)
    return raw


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    widths = table_widths(ncols, rows)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = doc.styles["Scientific Table"]
            add_inline(p, value)
            if r_idx == 0:
                shade(cell, LIGHT)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(INK)
        set_cant_split(table.rows[r_idx])
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    def cells(line: str) -> list[str]:
        return [x.strip() for x in line.strip().strip("|").split("|")]

    rows = [cells(lines[start])]
    i = start + 2  # skip separator
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append(cells(lines[i]))
        i += 1
    return rows, i


def build_content(doc: Document, markdown: str) -> None:
    bullet_id = add_numbering(doc, "bullet")
    number_id = add_numbering(doc, "decimal")
    lines = markdown.splitlines()
    i = 0
    current_h1 = ""
    # References may follow the final STAR Methods paragraph on the same page.
    # Forcing a break here created a two-line orphan page in the frozen manuscript.
    page_break_sections = {"Introduction", "STAR Methods", "Figure legends", "Supplemental information"}
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        hm = re.match(r"^(#{2,5})\s+(.+)$", line)
        if hm:
            level = len(hm.group(1)) - 1
            title = hm.group(2).strip()
            if level == 1:
                current_h1 = title
            p = doc.add_paragraph(style=f"Heading {level}")
            if level == 1 and title in page_break_sections:
                p.paragraph_format.page_break_before = True
            add_inline(p, title)
            i += 1
            continue
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue
        bm = re.match(r"^\s*-\s+(.+)$", line)
        if bm:
            p = doc.add_paragraph(style="Bullet Body")
            apply_num(p, bullet_id)
            add_inline(p, bm.group(1))
            i += 1
            continue
        nm = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if nm and current_h1 != "References":
            p = doc.add_paragraph(style="Number Body")
            apply_num(p, number_id)
            add_inline(p, nm.group(2))
            i += 1
            continue

        chunks = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip():
            nxt = lines[i]
            if re.match(r"^#{2,5}\s+", nxt) or re.match(r"^\s*-\s+", nxt) or re.match(r"^\s*\d+\.\s+", nxt):
                break
            if nxt.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
                break
            chunks.append(nxt.strip())
            i += 1
        text = " ".join(chunks)
        if current_h1 == "References":
            p = doc.add_paragraph(style="Reference Entry")
        elif current_h1 in {"Figure legends", "Supplemental information"}:
            p = doc.add_paragraph(style="Figure Legend")
        else:
            p = doc.add_paragraph(style="Normal")
        # Fully justified lines containing long, unbreakable checksums create
        # publication-distracting word spacing. Keep provenance paragraphs
        # left aligned while preserving the journal body style elsewhere.
        if "SHA-256" in text:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, text)


def main() -> None:
    markdown = INPUT.read_text(encoding="utf-8")
    title_match = re.match(r"^#\s+(.+)$", markdown, flags=re.M)
    if not title_match:
        raise ValueError("Assembled Markdown lacks title")
    title = title_match.group(1).strip()
    start = markdown.index("## Summary")
    content = markdown[start:]

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    doc.core_properties.title = title
    doc.core_properties.subject = "Public-data analysis of nintedanib-associated liver injury in fibrotic interstitial lung disease"
    doc.core_properties.author = "Study authors (administrative fields pending author verification)"
    doc.core_properties.keywords = "nintedanib, interstitial lung disease, DILI, proteomics, pharmacovigilance, virtual perturbation"
    add_title_page(doc, title)
    build_content(doc, content)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
